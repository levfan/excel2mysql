"""Win7 线（cli.py + core）用例集：一台机器上同时给 Python 3.8 和 3.12 跑。

为什么要在两个栈上跑同一份：Win7 线唯一的风险就是"旧栈上 pandas/numpy/解释器行为不一样"，
同一份断言、同一批输入，两边都必须给出同样的结果才算"core 零改动"这句话成立。
只在一边跑绿是证明不了这件事的。

覆盖：
  U 旧栈兼容守卫（AST 扫 3.9+ 语法与 stdlib、依赖版本契约、cli 不含 Qt 引用）
  V core 跨栈输出契约（固定输入 → 写死的期望表名/字段/类型/主键/DDL）
  W CLI 参数与连接配置（配置文件 utf-8/GBK、环境变量与命令行优先级、非法输入）
  X CLI 端到端·假连接（建表与插入调用序列、回滚、前缀、dry-run、退出码、日志形态）
  Y CLI 输入收集（目录递归、压缩包展开、不支持类型、路径不存在）
  G MySQL 真实读写（无凭据时显式记为 SKIPPED，绝不冒充通过）

运行：
  旧栈  .venv38\\Scripts\\python  pack_test38.py      （Win7 打包闸门用这个）
  新栈  .venv\\Scripts\\python    pack_test38.py      （证明两栈结果一致）
  冻结  dist_packtest38\\packtest38\\packtest38.exe
要连库真跑（G 组），先在环境里备好凭据：
  set PT_HOST=127.0.0.1& set PT_PORT=3306& set PT_USER=root& set PT_PASS=密码& set PT_DB=测试库
"""
import ast
import contextlib
import io
import locale
import os
import re
import shutil
import sys
import tempfile
import traceback
import zipfile

import pandas as pd

RESULTS = {"pass": 0, "fail": 0, "skip": 0}
FAILURES = []
SKIPPED = []

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass


def check(label, cond, detail=""):
    ok = bool(cond)
    RESULTS["pass" if ok else "fail"] += 1
    if not ok:
        FAILURES.append(label)
    print(("  OK   " if ok else "  FAIL ") + label + ("  | " + str(detail) if detail else ""))
    return ok


def skip(label, why):
    RESULTS["skip"] += 1
    SKIPPED.append((label, why))
    print(f"  SKIP {label}  | {why}")


def section(title):
    print(f"\n----- {title} -----")


SUMMARY_PRINTED = False


def _crash_report():
    if not SUMMARY_PRINTED:
        print("\n" + "=" * 72)
        print(f"用例脚本异常中止：已跑 通过 {RESULTS['pass']}  失败 {RESULTS['fail']}  "
              f"跳过 {RESULTS['skip']}")
        for f in FAILURES:
            print("  - " + f)
        print("=" * 72)


import atexit
atexit.register(_crash_report)

ROOT = os.path.dirname(os.path.abspath(__file__))
IS_OLD = sys.version_info[:2] == (3, 8)
STACK = "3.8/Win7 线" if IS_OLD else ("%d.%d 线" % sys.version_info[:2])
_tmpdir = tempfile.mkdtemp(prefix="packtest38_")


def source_of(mod_name):
    """取模块源码用于静态守卫；冻结态包里没有 .py，返回 None 由调用方记 SKIP。"""
    p = os.path.join(ROOT, mod_name)
    return open(p, encoding="utf-8").read() if os.path.isfile(p) else None


# ---------------------------------------------------------------------------
section("U 旧栈兼容守卫（这些语法/依赖在 Python 3.8 上会直接炸）")

import cli as cli_mod
import core.infer
import core.merge_check
import core.mysql_io
import core.naming
import core.sql_script
import core.zip_import

check("U1 当前解释器与用例集预期一致",
      True, "Python %s（%s）· pandas=%s numpy=%s"
      % (sys.version.split()[0], STACK, pd.__version__,
         __import__("numpy").__version__))

# 冻结态里 cli.py 没有源码，静态守卫逐条 SKIP，不参与通过数
_GUARD_TARGETS = ["cli.py", os.path.join("core", "infer.py"), os.path.join("core", "mysql_io.py"),
                  os.path.join("core", "sql_script.py"), os.path.join("core", "naming.py"),
                  os.path.join("core", "zip_import.py"), os.path.join("core", "merge_check.py"),
                  os.path.join("core", "log_text.py")]
_sources = {t: source_of(t) for t in _GUARD_TARGETS}

if any(v is None for v in _sources.values()):
    skip("U2 AST 扫描 3.9+ 语法与 stdlib", "冻结态包内没有 .py 源码，静态守卫只在开发态生效")
    skip("U3 无 PEP 604 联合类型注解", "同上")
    skip("U4 无 3.9+ 专属 stdlib 与 API", "同上")
else:
    # PEP 604（X | None）与 PEP 585（list[str]）写在注解里会在 import 时求值，3.8 直接 TypeError。
    # compileall 抓不到——它是合法语法；只有真 import 或按 AST 检查才暴露。
    _pep604, _pep585, _banned_import, _banned_api = [], [], [], []
    _BANNED_MODULES = {"zoneinfo", "graphlib", "tomllib"}
    _BANNED_ATTRS = {"removeprefix", "removesuffix", "batch_size"}
    _BANNED_FUNCTOOLS = {"cache"}
    for name, src in _sources.items():
        tree = ast.parse(src)
        for node in ast.walk(tree):
            anns = []
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                a = node.args
                anns = [x.annotation for x in list(a.args) + list(a.kwonlyargs)]
                if a.vararg:
                    anns.append(a.vararg.annotation)
                if a.kwarg:
                    anns.append(a.kwarg.annotation)
                anns.append(node.returns)
            elif isinstance(node, ast.AnnAssign):
                anns = [node.annotation]
            for an in anns:
                if an is None:
                    continue
                for sub in ast.walk(an):
                    if isinstance(sub, ast.BinOp) and isinstance(sub.op, ast.BitOr):
                        _pep604.append("%s:%s" % (name, getattr(an, "lineno", "?")))
                    if (isinstance(sub, ast.Subscript) and isinstance(sub.value, ast.Name)
                            and sub.value.id in ("list", "dict", "tuple", "set", "frozenset")):
                        _pep585.append("%s:%s" % (name, getattr(an, "lineno", "?")))
            if isinstance(node, ast.Import):
                for al in node.names:
                    if al.name.split(".")[0] in _BANNED_MODULES:
                        _banned_import.append("%s:%s %s" % (name, node.lineno, al.name))
            elif isinstance(node, ast.ImportFrom):
                if (node.module or "").split(".")[0] in _BANNED_MODULES:
                    _banned_import.append("%s:%s %s" % (name, node.lineno, node.module))
            elif isinstance(node, ast.Attribute):
                if node.attr in _BANNED_ATTRS:
                    _banned_api.append("%s:%s .%s()" % (name, node.lineno, node.attr))
                if (node.attr in _BANNED_FUNCTOOLS and isinstance(node.value, ast.Name)
                        and node.value.id in ("functools", "lru_cache")):
                    _banned_api.append("%s:%s functools.%s" % (name, node.lineno, node.attr))
    check("U2 被守卫的模块全部可 AST 解析", True, "、".join(os.path.basename(t) for t in _GUARD_TARGETS))
    check("U3 无 PEP 604 联合类型注解（X | None）", not _pep604,
          "; ".join(sorted(set(_pep604))) or "干净")
    check("U4 无 PEP 585 内置泛型注解（list[str] 等）", not _pep585,
          "; ".join(sorted(set(_pep585))) or "干净")
    check("U5 无 3.9+ 专属 stdlib / API", not (_banned_import + _banned_api),
          "; ".join(sorted(set(_banned_import + _banned_api))) or "干净")

cli_src = _sources.get("cli.py")
if cli_src is None:
    skip("U6 cli.py 不引用任何 Qt", "冻结态没有 cli.py 源码")
else:
    check("U6 cli.py 不引用任何 Qt",
          not re.search(r"\bPySide|PyQt|\bQt(Core|Widgets|Gui)\b", cli_src), "无 Qt 依赖才能进 Win7 包")
    check("U7 cli.py 不 import main（避免把 GUI 拖进旧包）",
          not re.search(r"^\s*(?:import|from)\s+main\b", cli_src, re.M))
    # 报错展开已经收进 core/log_text.py 三条线共用（cli 与 GUI 各一份必然有一天只改一处），
    # 所以这条断的是两件事：共用那份里真的写着错误码展开，cli 真的在用它而不是自己另写。
    _log_src = _sources.get(os.path.join("core", "log_text.py")) or ""
    check("U8 数据库报错展开在 core/log_text 里，cli 是引用它而不是另抄一份",
          "MySQL 错误码" in _log_src and "args[0]" in _log_src
          and re.search(r"^\s*from core\.log_text import .*describe_db_error", cli_src, re.M)
          is not None and "def describe_db_error" not in cli_src,
          "core/log_text.py 命中=%s，cli 引用=%s，cli 自带定义=%s"
          % ("MySQL 错误码" in _log_src,
             bool(re.search(r"^\s*from core\.log_text import", cli_src, re.M)),
             "def describe_db_error" in cli_src))

# 运行时取证：把 cli 及其依赖全 import 之后，进程里不该有任何 Qt 模块。
# 这条在冻结态照样成立，而且比 grep 源码硬——真被加载了才算数。
_qt_loaded = sorted(m for m in sys.modules
                    if m.split(".")[0] in ("PySide6", "PySide2", "PyQt5", "PyQt6", "shiboken6"))
check("U8b cli 加载后进程里没有 Qt 模块（运行时取证，冻结态同样有效）",
      not _qt_loaded, _qt_loaded or "无")
check("U8c core + cli 全链路都在当前栈上 import 成功（旧栈能 import 就是 3.8 兼容的硬证据）",
      all(m is not None for m in (cli_mod, core.infer, core.mysql_io, core.sql_script,
                                  core.naming, core.zip_import, core.merge_check,
                                  core.log_text)),
      "Python %s" % sys.version.split()[0])

def installed_version(dist):
    """取"装进环境里的那个版本"，靠发行版元数据而不是模块自己的 __version__。

    PyMySQL 1.1.2 把 `pymysql.__version__` 硬编码成了 "1.4.6"（上游遗留 bug），
    拿它跟 requirements 的 pin 比对会永久误报不一致；`importlib.metadata.version`
    读的是 dist-info，才是pip 实际装上的版本。
    """
    try:
        import importlib.metadata as md
    except ImportError:
        return None
    try:
        return md.version(dist)
    except Exception:
        return None


_V = {n: installed_version(n) for n in ("numpy", "pandas", "PyMySQL", "openpyxl",
                                        "xlrd", "pypinyin", "tzdata", "pyinstaller",
                                        "cryptography")}

if IS_OLD:
    check("U9 旧栈依赖的发行版元数据全部可读",
          all(_V.values()), " ".join("%s=%s" % (k, v or "?") for k, v in _V.items()))
    check("U10 旧栈里 numpy 仍是 1.24.4（1.25 起要 Python>=3.9）",
          _V["numpy"] == "1.24.4", _V["numpy"])
    check("U11 旧栈里 pandas 仍是 2.0.3（2.1 起要 Python>=3.9）",
          _V["pandas"] == "2.0.3", _V["pandas"])
    check("U11b PyMySQL 版本以 dist-info 为准（模块自报的 __version__ 是上游 bug）",
          _V["PyMySQL"] == "1.1.2" and __import__("pymysql").__version__ != "1.1.2",
          "dist-info=1.1.2, pymysql.__version__=%s" % __import__("pymysql").__version__)
    # MySQL 8 的 caching_sha2_password 在服务端缓存失效后的第一个连接要走 full auth，
    # 那一步要 RSA；pymysql 在 import 时就把有没有 cryptography 算成了这个开关，
    # 所以缺包不会让导入报错，只会在现场连不上库——必须在这里钉住。
    import pymysql._auth as _auth
    check("U11c MySQL 8 冷缓存首连要 RSA，cryptography 必须在场",
          _auth._have_cryptography,
          "pymysql._auth._have_cryptography=%s" % _auth._have_cryptography)
else:
    skip("U9-U11b requirements38 版本契约", "本轮跑在 %s，版本契约由 .venv38 那一轮核" % STACK)

req = source_of("requirements38.txt")
if req is None:
    skip("U12 requirements38.txt 与打包脚本存在", "冻结态没有 requirements 文件")
else:
    for pin in ("numpy==1.24.4", "pandas==2.0.3", "PyMySQL==1.1.2",
                "openpyxl==3.1.5", "xlrd==2.0.2", "pypinyin==0.55.0",
                "tzdata==2026.5", "pyinstaller==5.13.2", "cryptography==47.0.0"):
        check("U13 requirements38 钉住 %s" % pin, pin in req)
    check("U14 requirements38 写明了「这是最后能装进 3.8 的版本」这件事",
          "Windows 7" in req and "6.x" in req and "requires Python >= 3.9" in req)
    if IS_OLD:
        # 直接按 requirements 语法解析（name==version），不走 configparser：
        # configparser 会把选项名统一转小写，PyMySQL 变成 pymysql 后比对起来更绕
        _pins = dict(re.findall(r"^([A-Za-z0-9][A-Za-z0-9_.\-]*)==([0-9A-Za-z_.\-]+)[ \t]*$",
                                req, re.M))
        check("U15 requirements38 里每条依赖都钉死了版本",
              len(_pins) >= 8 and not re.search(
                  r"^[A-Za-z][A-Za-z0-9_.\-]*(==|>=|<=|~=)?\s*$", req, re.M), sorted(_pins))
        _mismatch, _unknown = [], []
        for name, want in sorted(_pins.items()):
            got = _V.get(name)
            if got is None:
                _unknown.append(name)
            elif got != want:
                _mismatch.append("%s 实装 %s != 钉住 %s" % (name, got, want))
        check("U16b 旧栈实际装的与 requirements38 写的一致", not _mismatch,
              "; ".join(_mismatch) or "一致")
        check("U16c 每个 pin 都查到了发行版元数据", not _unknown,
              "查不到 " + "、".join(_unknown) if _unknown else "全部命中")
    else:
        skip("U15-U16c 旧栈实际装的与 requirements38 写的一致",
             "本轮跑在 %s，装的是 GUI 线依赖" % STACK)

if IS_OLD:
    check("U17 旧栈 PyInstaller 是 5.x（6.x 只支持 Windows 8+，bootloader 不为 Win7 构建）",
          (_V.get("pyinstaller") or "").startswith("5."), _V.get("pyinstaller"))
else:
    skip("U17 PyInstaller 5.x 守卫", "PyInstaller 装在 .venv38，本轮是 %s" % STACK)

# requirements 文件必须纯 ASCII：Python 3.8 出厂的 pip 21.1.1 用
# locale.getpreferredencoding() 解 requirements（zh-CN Windows 上是 cp936），
# 带 UTF-8 中文注释直接 UnicodeDecodeError 装不上依赖。新版 pip 只是碰巧容忍。
for _rf in ("requirements38.txt", "requirements.txt"):
    _raw = None
    _p = os.path.join(ROOT, _rf)
    if os.path.isfile(_p):
        with open(_p, "rb") as _fh:
            _raw = _fh.read()
    if _raw is None:
        skip("U18 %s 全 ASCII（老 pip 才解得开）" % _rf, "冻结态包内没有该文件")
        skip("U19 %s 能按本机首选码页解码" % _rf, "同上")
        continue
    _non_ascii = sorted({b for b in bytearray(_raw) if b > 127})
    check("U18 %s 全 ASCII（老 pip 才解得开）" % _rf, not _non_ascii,
          "出现 %d 个非 ASCII 字节" % len(_non_ascii) if _non_ascii else "干净")
    try:
        _raw.decode(locale.getpreferredencoding(False))
        check("U19 %s 能按本机首选码页解码" % _rf, True,
              locale.getpreferredencoding(False))
    except UnicodeDecodeError as e:
        check("U19 %s 能按本机首选码页解码" % _rf, False, "%s → %s"
              % (locale.getpreferredencoding(False), e))
    check("U20 %s 的依赖行都能被 pip 语法接受（name[extras]=ver）" % _rf,
          all(re.match(r"^[A-Za-z0-9_.\-]+(\[[A-Za-z0-9_,\-]+\])?(==|>=|<=|~=)?[A-Za-z0-9_.\-]*$",
                       ln.strip())
              for ln in _raw.decode("ascii", errors="replace").splitlines()
              if ln.strip() and not ln.strip().startswith("#")))


# ---------------------------------------------------------------------------
section("V core 跨栈输出契约（固定输入 → 两边必须给出同样的表名/类型/DDL）")


def make_xlsx(name, sheets, folder=None):
    """sheets: 单个 DataFrame（写到 Sheet1）或 {sheet名: DataFrame}；返回写出的路径。

    不能只收 dict：DataFrame 也有 .items()（列名→Series），传错了会"看起来跑了"
    但每个列被当成一个 Sheet 写出去，后面的断言全部对不上——这类假失败最难查。
    """
    p = os.path.join(folder or _tmpdir, name)
    if isinstance(sheets, pd.DataFrame):
        sheets = {"Sheet1": sheets}
    with pd.ExcelWriter(p, engine="openpyxl") as xw:
        for sn, df in sheets.items():
            df.to_excel(xw, sheet_name=sn, index=False)
    return p


from core.infer import build_table_def, make_unique_names
from core.mysql_io import build_ddl

v_orders = make_xlsx("订单（ods_t_order）.xlsx", pd.DataFrame({
    "订单编号": ["SO0001", "SO0002", "SO0003"],
    "客户名称": ["甲公司", "乙公司", "丙公司"],
    "金额": [12.5, 3.25, 100.0],
    "下单日期": ["2026-01-05", "2026-02-06", "2026-03-07"],
    "数量": [3, 7, 1],
}))
t = build_table_def(v_orders)
cmt = {c["comment"]: c["type"] for c in t["columns"]}
check("V1 括号内英文作表名", t["table_name"] == "ods_t_order", t["table_name"])
check("V2 括号前中文作表注释", t["table_comment"] == "订单", t["table_comment"])
check("V3 表头逐音节转拼音蛇形命名",
      [c["field"] for c in t["columns"]]
      == ["ding_dan_bian_hao", "ke_hu_ming_cheng", "jin_e", "xia_dan_ri_qi", "shu_liang"],
      [c["field"] for c in t["columns"]])
check("V4 编号列→VARCHAR（不是 INT）", cmt["订单编号"].startswith("VARCHAR"), cmt["订单编号"])
check("V5 金额列→DECIMAL(16,2)", cmt["金额"] == "DECIMAL(16,2)", cmt["金额"])
check("V6 日期串→DATE", cmt["下单日期"] == "DATE", cmt["下单日期"])
check("V7 整数列→INT", cmt["数量"] == "INT", cmt["数量"])
check("V8 主键选中的是编号列", t["primary_key"] == "ding_dan_bian_hao", t.get("primary_key"))
check("V9 数据行数不随栈变化", len(t["df"]) == 3, len(t["df"]))
ddl = build_ddl(t["table_name"], t["table_comment"], t["columns"])
check("V10 DDL 文本逐字符稳定（utf8mb4_general_ci / InnoDB）",
      ddl.startswith("CREATE TABLE `ods_t_order` (")
      and "ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_general_ci" in ddl
      and "PRIMARY KEY (`ding_dan_bian_hao`)" in ddl, ddl.replace("\n", "\\n")[:120])
check("V11 主键列被写成 NOT NULL",
      "`ding_dan_bian_hao` VARCHAR(64) NOT NULL" in ddl)

v_dec = make_xlsx("库容（dim_rong）.xlsx", pd.DataFrame({
    "编码": ["A1", "A2", "A3"],
    "总库容（亿立方米）": [4.375, 4.160, 2.930],
}))
td = build_table_def(v_dec)
tc = {c["comment"]: c["type"] for c in td["columns"]}
check("V12 三位小数按实测精度出 DECIMAL(7,3)（不是 20,6 也不是圆到 2 位）",
      tc["总库容（亿立方米）"] == "DECIMAL(7,3)", tc["总库容（亿立方米）"])

v_pure = make_xlsx("纯数字（7）.xlsx", pd.DataFrame({"编号": [1, 2]}))
tp = build_table_def(v_pure)
# 期望值抄主用例集 S3 的权威口径：括号内纯数字是序号，回退拼音后**保留**序号
check("V13 括号内纯数字是序号不是表名（回退拼音并保留序号）",
      tp["table_name"] == "chun_shu_zi_7", tp["table_name"])

v_cn = make_xlsx("说明（人事处人员）.xlsx", pd.DataFrame({"编号": [1, 2]}))
check("V14 括号内汉字不作表名（避免汉字表名）",
      build_table_def(v_cn)["table_name"] == "shuo_ming_ren_shi_chu_ren_yuan",
      build_table_def(v_cn)["table_name"])

v_ver = make_xlsx("订单表（V2）.xlsx", pd.DataFrame({"编号": [1, 2]}))
tv = build_table_def(v_ver)
check("V14b 括号内版本号不作表名：回退拼音保留版本段，首字母+大写为 DDBV2",
      tv["table_name"] == "ding_dan_biao_v2"
      and tv["table_comment"] == "订单表（V2）"
      and build_table_def(v_ver, initials=True, upper=True)["table_name"] == "DDBV2",
      (tv["table_name"], tv["table_comment"],
       build_table_def(v_ver, initials=True, upper=True)["table_name"]))

dup_a = make_xlsx("同名表.xlsx", pd.DataFrame({"编号": [1]}))
dup_b = os.path.join(_tmpdir, "dup", "同名表.xlsx")
os.makedirs(os.path.dirname(dup_b), exist_ok=True)
shutil.copy(dup_a, dup_b)
tables = [build_table_def(dup_a), build_table_def(dup_b)]
make_unique_names(tables)
check("V15 撞名自动加 _2",
      [x["table_name"] for x in tables] == ["tong_ming_biao", "tong_ming_biao_2"],
      [x["table_name"] for x in tables])

from core.sql_script import (
    apply_table_prefix_sql, extract_create_tables, split_sql_statements,
    prepend_drop_tables,
)

_script = ("CREATE TABLE `ods_t_order` (`id` INT) COMMENT='订单';\n"
           "INSERT INTO `ods_t_order` VALUES (1);\n"
           "-- 注释里有 ; 分号\n"
           "CREATE TABLE `dim_a` (`x` VARCHAR(20) DEFAULT 'a;b');\n")
check("V16 SQL 脚本表名提取跨栈一致",
      [e["name"] for e in extract_create_tables(_script)] == ["ods_t_order", "dim_a"],
      extract_create_tables(_script))
_new, _ren, _skip = apply_table_prefix_sql(_script, "ads_")
check("V17 脚本加前缀的改写对象与处数跨栈一致",
      [(r["old"], r["new"], r["n"]) for r in _ren]
      == [("ods_t_order", "ads_ods_t_order", 2), ("dim_a", "ads_dim_a", 1)], _ren)
check("V17b 加前缀后脚本里两个表名都被改写",
      "`ads_ods_t_order`" in _new and "`ads_dim_a`" in _new and "`ods_t_order`" not in _new)
check("V17c 重复应用不叠加前缀",
      apply_table_prefix_sql(_new, "ads_")[2] == ["ads_ods_t_order", "ads_dim_a"],
      apply_table_prefix_sql(_new, "ads_")[2])
check("V18 字符串与注释里的分号不当成语句边界",
      len(split_sql_statements(_script)) == 3, split_sql_statements(_script))
_drop_text, _drop_tabs = prepend_drop_tables(
    "-- CREATE TABLE ghost (id INT);\n"
    "CREATE TABLE `db`.`t1` (id INT);\nCREATE TABLE t2 (id INT);\n"
    "CREATE TABLE t1 (id INT);\n")
check("V18b 覆盖执行只在整句 CREATE TABLE 前补 DROP，倒序去重且保留 schema 限定名",
      _drop_tabs == ["`db`.`t1`", "t2", "t1"]
      and _drop_text.startswith("DROP TABLE IF EXISTS t1;\n"
                                "DROP TABLE IF EXISTS t2;\n"
                                "DROP TABLE IF EXISTS `db`.`t1`;\n\n"),
      (_drop_tabs, _drop_text.splitlines()[:3]))
check("V18c 没有建表语句时不补 DROP",
      prepend_drop_tables("SELECT 1;") == ("SELECT 1;", []))

from core.merge_check import MergedCellsError, ensure_no_merged_cells

merged_path = os.path.join(_tmpdir, "含合并.xlsx")
from openpyxl import Workbook

_wb = Workbook()
_ws = _wb.active
_ws.append(["姓名", "分数"])
_ws.append(["张三", 90])
_ws.merge_cells("A2:A3")
_ws.append(["", 80])
_wb.save(merged_path)
try:
    build_table_def(merged_path)
    check("V19 合并单元格文件被拒绝导入", False, "竟然解析成功了")
except MergedCellsError as e:
    check("V19 合并单元格文件被拒绝导入", True, "合并区 %d 处" % len(e.merged))
except Exception as e:
    check("V19 合并单元格文件被拒绝导入", False, "抛的是 %s: %s" % (type(e).__name__, e))

merged_two = os.path.join(_tmpdir, "两行合并.xlsx")


def _mk_raw_header(path, extra_merge=None, other_sheet=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "原始表"
    ws.append(["序号", "技术公司", None, "财务公司", None])
    ws.append([None, "小计", "高层", "小计", "高层"])
    for box in ("A1:A2", "B1:C1", "D1:E1"):
        ws.merge_cells(box)
    ws.append([1, "甲", 10, 20, 30])
    ws.append([2, "乙", 11, 21, 31])
    if extra_merge:
        ws.merge_cells(extra_merge)
    if other_sheet:
        d = wb.create_sheet("其他")
        d.append(["a", "b"])
        d.append([1, 2])
        d.merge_cells("A1:B1")
    wb.save(path)


_mk_raw_header(merged_two)
_t = build_table_def(merged_two)
check("V19b 两行合并表头压成「父-子」一行",
      [c["comment"] for c in _t["columns"]]
      == ["序号", "技术公司-小计", "技术公司-高层", "财务公司-小计", "财务公司-高层"],
      [c["comment"] for c in _t["columns"]])
check("V19c 展开后字段名、数据行与主键照常产出",
      [c["field"] for c in _t["columns"]]
      == ["xu_hao", "ji_shu_gong_si_xiao_ji", "ji_shu_gong_si_gao_ceng",
          "cai_wu_gong_si_xiao_ji", "cai_wu_gong_si_gao_ceng"]
      and list(_t["df"].iloc[0]) == [1, "甲", 10, 20, 30]
      and _t["primary_key"] == "xu_hao"
      and _t["merged_header"] == {"sheet": "原始表", "depth": 2},
      ([c["field"] for c in _t["columns"]], list(_t["df"].iloc[0]),
       _t["primary_key"], _t["merged_header"]))
merged_one_row = os.path.join(_tmpdir, "单行合并.xlsx")
_wb2 = Workbook()
_ws2 = _wb2.active
_ws2.append(["组A", None, "姓名"])
_ws2.append(["甲", "乙", "丙"])
_ws2.merge_cells("A1:B1")
_wb2.save(merged_one_row)
merged_other = os.path.join(_tmpdir, "其他页合并.xlsx")
_mk_raw_header(merged_other, other_sheet=True)


def _mk_raw_shape(path, top, second, merges, third=None, data=(1, 10, 20)):
    wb = Workbook()
    ws = wb.active
    ws.title = "原始表"
    ws.append(top)
    ws.append(second)
    if third is not None:
        ws.append(third)
    for box in merges:
        ws.merge_cells(box)
    ws.append(list(data))
    wb.save(path)
    return path


merged_three = _mk_raw_shape(
    os.path.join(_tmpdir, "三行合并表头.xlsx"),
    ["序号", "技术公司", None], [None, "华北", None],
    ("A1:A3", "B1:C1", "B2:C2"), third=[None, "小计", "高层"], data=[1, 10, 20])
merged_child = _mk_raw_shape(
    os.path.join(_tmpdir, "子行合并表头.xlsx"),
    ["序号", "技术公司", None, "财务公司", None],
    [None, "小计", "高层", "小计", "高层"],
    ("A1:A2", "B1:C1", "D1:E1", "B2:C2"), data=[1, 10, 20, 30, 40])
merged_block = _mk_raw_shape(
    os.path.join(_tmpdir, "2x2合并表头.xlsx"),
    ["序号", "技术公司", None, "财务公司"],
    [None, "小计", "高层", "小计"],
    ("A1:A2", "B1:C2", "D1:D2"), data=[1, 10, 20, 30])
for _label, _path in (("单行合并表头", merged_one_row), ("其他页也合并", merged_other),
                      ("三行合并表头", merged_three), ("第 2 行子合并", merged_child),
                      ("2x2 合并块", merged_block)):
    try:
        build_table_def(_path)
        _ok = False
    except MergedCellsError:
        _ok = True
    check("V19d 仍拒绝：" + _label, _ok)

from core.zip_import import ZipArchiveError, count_excel_entries, extract_excel_archive

zip_src = os.path.join(_tmpdir, "inner", "人员.xlsx")
os.makedirs(os.path.dirname(zip_src), exist_ok=True)
make_xlsx("人员.xlsx", pd.DataFrame({"工号": ["E1", "E2"], "姓名": ["张三", "李四"]}),
          folder=os.path.dirname(zip_src))
zip_path = os.path.join(_tmpdir, "数据包.zip")
with zipfile.ZipFile(zip_path, "w") as zf:
    zf.write(zip_src, "包内/人员.xlsx")
    zf.writestr("包内/说明.txt", "非 Excel")
    zf.writestr("包内/嵌套.zip", "假 zip")
check("V20 预扫包内 Excel 条目数", count_excel_entries(zip_path) == 1, count_excel_entries(zip_path))
r = extract_excel_archive(zip_path, os.path.join(_tmpdir, "out"))
check("V21 展开只取 Excel、忽略非 Excel、跳过嵌套 zip",
      len(r["files"]) == 1 and len(r["ignored"]) == 1 and r["nested_zips"] == 1,
      "files=%d ignored=%d nested=%d" % (len(r["files"]), len(r["ignored"]), r["nested_zips"]))
check("V22 包内中文名保留（表名靠它）",
      r["files"][0]["origin"].replace("\\", "/") == "包内/人员.xlsx", r["files"][0]["origin"])
tz = os.path.join(_tmpdir, "trav.zip")
with zipfile.ZipFile(tz, "w") as zf:
    zf.writestr("../../evil.xlsx", "x")
r2 = extract_excel_archive(tz, os.path.join(_tmpdir, "out2"))
check("V23 zip slip 路径被净化（不落 ../）",
      all("..".split(os.sep)[0] not in os.path.abspath(x["path"]).replace(os.path.abspath(
          os.path.join(_tmpdir, "out2")), "") for x in r2["files"]) or not r2["files"],
      [x["path"] for x in r2["files"]])
with open(os.path.join(_tmpdir, "notzip.zip"), "wb") as fh:
    fh.write(b"plain text, not a zip")
try:
    extract_excel_archive(os.path.join(_tmpdir, "notzip.zip"), os.path.join(_tmpdir, "out3"))
    check("V24 坏压缩包抛 ZipArchiveError", False, "没抛")
except ZipArchiveError as e:
    check("V24 坏压缩包抛 ZipArchiveError", True, str(e))

xls_env = os.environ.get("PT_XLS")
if xls_env and os.path.isfile(xls_env):
    tx = build_table_def(xls_env)
    check("V25 真 .xls 可读（xlrd）", bool(tx["columns"]), tx["table_name"])
else:
    skip("V25 真 .xls 可读", "未设 PT_XLS 指向真实 .xls 样本（本机没有该格式文件）")


# 「生成标准/简易 excel 文件」这条产物线：3.8 冻结态里 openpyxl 必须真能写，
# 写出来的标准件还要能被本线原样读回来（用户拿它去继续导入才有意义）
from core import meta_export as _mx
from openpyxl import load_workbook as _load

_vt = build_table_def(v_orders)
_std_root = os.path.join(_tmpdir, "export_std")
os.makedirs(_std_root, exist_ok=True)
_vtgt, _vwritten = _mx.export_tables([_vt], _std_root, True)
check("V26 标准件写出一个文件，命名是表注释（表名）",
      len(_vwritten) == 1 and os.path.basename(_vwritten[0]) == u"订单（ods_t_order）.xlsx",
      [os.path.basename(x) for x in _vwritten])
check("V27 标准件目录名是 meta_data_file_年月日",
      os.path.basename(_vtgt).startswith("meta_data_file_"), _vtgt)
if _vwritten:
    _wb = _load(_vwritten[0], read_only=True)
    _names = list(_wb.sheetnames)
    _wb.close()
    check("V28 标准件是元数据+数据两个 Sheet", _names == [u"元数据", u"数据"], _names)
    _back = build_table_def(_vwritten[0])
    check("V29 标准件回读走结构定义那条路", _back["source"] == "structured",
          _back["source"])
    check("V30 回读后字段名与类型一字不差",
          [(c["field"], c["type"]) for c in _back["columns"]]
          == [(c["field"], c["type"]) for c in _vt["columns"]],
          [(c["field"], c["type"]) for c in _back["columns"]])
    check("V31 回读后主键没丢",
          [c["field"] for c in _back["columns"] if c.get("primary_key")]
          == [c["field"] for c in _vt["columns"] if c.get("primary_key")],
          _back.get("primary_key"))
    check("V32 回读后行数不少", len(_back["df"].index) == len(_vt["df"].index),
          "%d -> %d" % (len(_vt["df"].index), len(_back["df"].index)))
_sim_root = os.path.join(_tmpdir, "export_sim")
os.makedirs(_sim_root, exist_ok=True)
_vtgt2, _vwritten = _mx.export_tables([_vt], _sim_root, False)
check("V33 简易件目录名是 simple_meta_data_file_年月日",
      os.path.basename(_vtgt2).startswith("simple_meta_data_file_"), _vtgt2)
if _vwritten:
    _wb2 = _load(_vwritten[0], read_only=True)
    _names2 = list(_wb2.sheetnames)
    _wb2.close()
    check("V34 简易件只有一个数据 Sheet", _names2 == [u"数据"], _names2)


# ---------------------------------------------------------------------------
section("W CLI 参数与连接配置")
@contextlib.contextmanager
def clean_env(**extra):
    """CLI 的用例必须在受控环境变量下跑：本机可能正巧设了 E2M_*/PT_*，会污染优先级断言。"""
    saved = {}
    keys = [k for k in list(os.environ) if k.startswith(("E2M_", "PT_"))]
    for k in keys:
        saved[k] = os.environ.pop(k)
    os.environ.update({k: v for k, v in extra.items() if v is not None})
    try:
        yield
    finally:
        for k in list(os.environ):
            if k.startswith(("E2M_", "PT_")):
                os.environ.pop(k, None)
        os.environ.update(saved)


_LOG_SEQ = [0]


def run_cli(argv, log_name=None, env=None):
    """跑一次 cli.run，返回 (退出码, 日志文件文本, 控制台文本)。

    日志一律落到 UTF-8 文件：Win7 线上控制台按 cp936 输出，重定向到管道时读出来是乱码，
    断言只有对着日志文件才靠得住。不传 log_name 也自动给一个，避免"断言读到了空串"
    这种假失败——第一版就在这里连红十几条。

    env 里的键在 scrub 之后写入，用来给 G 组传凭据（走环境变量，不进命令行也不落盘）。
    """
    args = list(argv)
    if log_name is None:
        _LOG_SEQ[0] += 1
        log_name = "auto%d" % _LOG_SEQ[0]
    log_path = os.path.join(_tmpdir, "logs", log_name + ".log")
    os.makedirs(os.path.dirname(log_path), exist_ok=True)
    if os.path.isfile(log_path):
        os.remove(log_path)
    args += ["--log", log_path]
    buf = io.StringIO()
    with clean_env(**(env or {})):
        with contextlib.redirect_stdout(buf):
            code = cli_mod.run(args)
    text = ""
    if os.path.isfile(log_path):
        with open(log_path, encoding="utf-8") as fh:
            text = fh.read()
    return code, text, buf.getvalue()


def log_lines(text):
    """日志文件每行前面带 [时间戳]，断言"行形状"时先剥掉它。

    不剥的话 `^[+x!-] \\[OK\\]` 这类正则在空列表上恒真，守卫会永远绿——
    第一版 Y11 就是这么假过的。
    """
    return [re.sub(r"^\[[0-9\-: ]+\] ", "", ln) for ln in text.splitlines() if ln.strip()]


code, _o, console = run_cli(["--help"])
check("W1 --help 走正常返回而不是把 SystemExit 抛出去", code == cli_mod.EXIT_OK, code)
check("W1b --help 里三种动作与配置项都列全了",
      all(k in console for k in ("--dry-run", "--check", "--config", "--prefix",
                                 "--swap-prefix", "--drop-if-exists", "--log")),
      console[:60])
check("W1c --help 给了退出码表（脚本要按码分流）",
      "退出码" in console and "1" in console, console[:60])

code, _o, console = run_cli(["--version"])
check("W1d --version 返回 0 并打印版本号",
      code == cli_mod.EXIT_OK and cli_mod.APP_VERSION in console, (code, console.strip()[:40]))

code, _o, console = run_cli(["--batch", "abc", "-d", "x"])
check("W1e 非法取值折成用法码 1（不能让 argparse 的 2 冒充「解析失败」）",
      code == cli_mod.EXIT_USAGE, (code, console.strip()[-60:]))

code, _o, console = run_cli(["--definitely-not-a-flag"])
check("W1f 未知参数同样回用法码 1", code == cli_mod.EXIT_USAGE, code)

code, out, _ = run_cli([])
check("W2 没有任何输入时报用法错误并退出码 1", code == cli_mod.EXIT_USAGE, code)
check("W3 用法错误在日志里说清缺什么", "没有输入" in out, out.strip().splitlines()[:1])

code, out, _ = run_cli(["--dry-run", "--check", _tmpdir])
check("W4 --dry-run 与 --check 互斥", code == cli_mod.EXIT_USAGE, code)
check("W5 互斥提示点出两个参数名", "--dry-run" in out and "--check" in out, out.strip()[:80])

# 收尾那行"退出码 N"是批处理排障第一眼看的字：早退路径（参数冲突、配置错误、
# 没有可处理输入、Ctrl+C）以前都报 0，而进程实际退 1/2/3/130 —— 现场会照着 0 继续往下跑
_logged = re.search(r"退出码\s+(\S+)", out)
check("W5b 日志里自报的退出码等于真实返回码（互斥路径）",
      bool(_logged) and _logged.group(1) == str(code),
      (_logged.group(1) if _logged else None, code))
code, out2, _ = run_cli([_tmpdir, "-d", "db1", "--prefix", "a_", "--swap-prefix", "b_",
                         "--dry-run"])
_l2 = re.search(r"退出码\s+(\S+)", out2)
check("W5c 前缀两开关同时给：退 1 且日志也写 1",
      code == cli_mod.EXIT_USAGE and _l2 and _l2.group(1) == "1",
      (code, _l2.group(1) if _l2 else None))
code, out3, _ = run_cli([os.path.join(_tmpdir, "根本不存在.xlsx"), "-d", "db1"])
_l3 = re.search(r"退出码\s+(\S+)", out3)
check("W5d 输入不存在时日志自报的码与返回值一致",
      _l3 is None or _l3.group(1) == str(code), (code, _l3.group(1) if _l3 else None))

code, out, _ = run_cli([os.path.join(_tmpdir, "订单（ods_t_order）.xlsx")])
check("W6 未指定目标库直接判用法错误（不做半截导入）", code == cli_mod.EXIT_USAGE, code)
check("W7 缺库提示给了三种补救办法",
      "--database" in out and "E2M_DB" in out and "--dry-run" in out, out.strip()[:120])

code, out, _ = run_cli([_tmpdir, "-d", "db1", "--prefix", "ods$", "--dry-run"])
check("W8 前缀含非法字符被拒绝", code == cli_mod.EXIT_USAGE, code)
check("W9 非法前缀提示允许的字符集", "只允许字母、数字、下划线" in out, out.strip()[:80])

code, out, _ = run_cli([_tmpdir, "-d", "db1", "--swap-prefix", "  ", "--dry-run"])
check("W10 --swap-prefix 清完为空判用法错误", code == cli_mod.EXIT_USAGE, code)

code, out, _ = run_cli([_tmpdir, "-d", "db1", "--prefix", "a_", "--swap-prefix", "b", "--dry-run"])
check("W11 --prefix 与 --swap-prefix 互斥", code == cli_mod.EXIT_USAGE, code)

code, out, _ = run_cli([_tmpdir, "-d", "db1", "--batch", "0", "--dry-run"])
check("W12 --batch 0 被拒绝（否则会死循环或整表不提交）", code == cli_mod.EXIT_USAGE, code)

cfg_path = os.path.join(_tmpdir, "conn_utf8.ini")
with open(cfg_path, "w", encoding="utf-8") as fh:
    fh.write("# 中文注释：内网测试库\n[mysql]\nhost = 10.0.0.9\nport = 3307\n"
             "user = reader\npassword = pw-from-file\ndatabase = filedemo\n")
cfg = cli_mod.read_config_file(cfg_path)
check("W13 配置文件读全五项",
      (cfg["host"], cfg["port"], cfg["user"], cfg["database"])
      == ("10.0.0.9", "3307", "reader", "filedemo"), cfg)
check("W14 配置文件里的中文注释不炸（UTF-8）", "内网测试库" in open(cfg_path, encoding="utf-8").read())

cfg_gbk = os.path.join(_tmpdir, "conn_gbk.ini")
with open(cfg_gbk, "wb") as fh:
    fh.write("# 记事本另存为 ANSI 的中文注释\n[mysql]\nhost = 10.0.0.8\nport = 3306\n"
             "user = gbk\npassword = x\ndatabase = gbkdemo\n".encode("gbk"))
cfg2 = cli_mod.read_config_file(cfg_gbk)
check("W15 GBK(记事本 ANSI) 配置文件退回 GBK 仍能读", cfg2["host"] == "10.0.0.8", cfg2)

try:
    cli_mod.read_config_file(os.path.join(_tmpdir, "不存在.ini"))
    check("W16 配置文件不存在时报可读的错误", False, "没抛")
except ValueError as e:
    check("W16 配置文件不存在时报可读的错误", "配置文件不存在" in str(e), e)

nosec = os.path.join(_tmpdir, "nosec.ini")
with open(nosec, "w", encoding="utf-8") as fh:
    fh.write("[other]\nhost = 1.2.3.4\n")
try:
    cli_mod.read_config_file(nosec)
    check("W17 缺 [mysql] 段时报错并提示段名", False, "没抛")
except ValueError as e:
    check("W17 缺 [mysql] 段时报错并提示段名", "mysql" in str(e), e)

check("W18 [connection] 段名也接受（与 README 口径一致）",
      cli_mod.CONFIG_SECTIONS == ("mysql", "connection"), cli_mod.CONFIG_SECTIONS)

# 优先级：命令行 > 环境变量 > 配置文件 > 默认
with clean_env(E2M_HOST="env-host", E2M_PORT="3399", E2M_DB="envdb", E2M_USER="envuser"):
    ns = cli_mod.build_parser().parse_args(["-c", cfg_path, "--host", "cli-host", "-d", "db1"])
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        merged = cli_mod.resolve_cfg(ns, cli_mod.Log(quiet=True))
check("W19 命令行覆盖环境变量与配置文件", merged["host"] == "cli-host", merged)
check("W20 环境变量覆盖配置文件（未被命令行覆盖的项）",
      merged["port"] == 3399 and merged["database"] == "db1", merged)
check("W21 配置文件提供命令行/环境都没给的项", merged["user"] == "envuser", merged)

with clean_env(E2M_PASS="env-secret"):
    ns = cli_mod.build_parser().parse_args(["-c", cfg_path])
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        merged = cli_mod.resolve_cfg(ns, cli_mod.Log(quiet=True))
check("W22 E2M_PASS 覆盖配置文件密码", merged["password"] == "env-secret", merged)
check("W23 日志里不回显密码明文", "env-secret" not in buf.getvalue() and "pw-from-file" not in buf.getvalue(),
      buf.getvalue().strip()[:100])

with clean_env():
    ns = cli_mod.build_parser().parse_args([])
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        merged = cli_mod.resolve_cfg(ns, cli_mod.Log(quiet=True))
check("W24 全默认值时 host/port/user 有兜底",
      (merged["host"], merged["port"], merged["user"]) == ("127.0.0.1", 3306, "root"), merged)

# --- W25-W31：随包带的默认连接配置（assets/conn.ini → 产物同级 conn.ini） ---
# 这四条断的是源文件本身，冻结态包内没有 assets/conn.ini，只能 SKIP：
# "产物里到底有没有带 conn.ini"由闸门按实际产物目录核（build_exe38.inspect_package），
# 在测试包里再收一份副本自证是自欺欺人。
W25_LABELS = ("W25 模板 assets/conn.ini 存在且能解析",
              "W25b 模板给的默认值合理（本机 3306、root，口令与库留空）",
              "W25c 模板里的键与解析器认的键完全对齐（改模板不会静默失效）",
              "W25d 模板不含任何真实口令（公开仓库里不能带）")
tpl_path = os.path.join(ROOT, "assets", "conn.ini")
if not os.path.isfile(tpl_path):
    for _lbl in W25_LABELS:
        skip(_lbl, "冻结态包内没有 assets/conn.ini 源文件，模板内容由开发态那轮核")
    tpl = {}
else:
    check(W25_LABELS[0], True, tpl_path)
    tpl = cli_mod.read_config_file(tpl_path)
    check(W25_LABELS[1],
          (tpl.get("host"), str(tpl.get("port")), tpl.get("user")) == ("127.0.0.1", "3306", "root")
          and not (tpl.get("password") or "").strip() and not (tpl.get("database") or "").strip(), tpl)
    check(W25_LABELS[2], set(tpl) == set(cli_mod.CONFIG_KEYS), sorted(tpl))
    check(W25_LABELS[3],
          "Root@" not in open(tpl_path, encoding="utf-8").read()
          and "password =" in open(tpl_path, encoding="utf-8").read())

cands = cli_mod.default_config_paths()
check("W26 自动发现顺序：先程序目录、再当前工作目录",
      len(cands) == 2 and os.path.basename(cands[0]) == "conn.ini"
      and os.path.isabs(cands[0]) and os.path.isabs(cands[1]), cands)

code, out, _ = run_cli([v_orders, "--dry-run"])
check("W26b 开发态（仓库根没有 conn.ini）不会误读 assets 模板，来源写的是内置默认值",
      code == cli_mod.EXIT_OK and "内置默认值" in out, out[:200])

auto_dir = os.path.join(_tmpdir, "autocfg")
shutil.rmtree(auto_dir, ignore_errors=True)
os.makedirs(auto_dir)
_real_cwd = os.getcwd()
with open(os.path.join(auto_dir, "conn.ini"), "w", encoding="utf-8") as fh:
    fh.write("[mysql]\nhost = auto-host\nport = 3399\nuser = auto-user\n"
             "password = auto-secret\ndatabase = autodb\n")
try:
    os.chdir(auto_dir)
    code, out, _ = run_cli([v_orders, "--dry-run"])
    check("W27 不带 -c 时自动读当前目录的 conn.ini 并在来源里标明「自动发现」",
          code == cli_mod.EXIT_OK and "（自动发现）" in out and "auto-host" in out, out[:260])
    check("W27b 自动发现的配置里没有回显口令值",
          "auto-secret" not in out, out[:200])
    code, out, _ = run_cli([v_orders, "--dry-run", "-c", cfg_path])
    check("W27c 显式 -c 优先于自动发现（来源里不再标自动发现，用的是指定的那个文件）",
          code == cli_mod.EXIT_OK and "（自动发现）" not in out and "10.0.0.9" in out, out[:200])
    # 注意用 run_cli 的 env 参数：它内部会先 scrub 掉所有 E2M_*/PT_* 再写入，
    # 从外面 set 的变量会被它自己的清场逻辑抹掉
    code, out, _ = run_cli([v_orders, "--dry-run"], env={"E2M_HOST": "env-beats-auto"})
    check("W27d 环境变量仍然压过自动发现的配置文件",
          "env-beats-auto" in out and "（自动发现）" in out, out[:220])
    os.remove(os.path.join(auto_dir, "conn.ini"))
    with open(os.path.join(auto_dir, "conn.ini"), "w", encoding="utf-8") as fh:
        fh.write("[wrong]\nhost = 1.2.3.4\n")
    code, out, _ = run_cli([v_orders, "--dry-run"])
    check("W28 自动发现的配置写坏了：报 WARN 并退回内置默认值，命令照常跑完",
          code == cli_mod.EXIT_OK and "默认配置忽略" in out and "127.0.0.1" in out, out[:260])
finally:
    os.chdir(_real_cwd)

# --- W29-W31d：随包分发的 README.md 与真实参数不许漂移 ---
# 项目里只剩这一份文档，它同时是发到现场手里的一手资料：写错一个开关名或一个退出码，
# 对面就会照着敲错的命令。所以用 argparse 与 cli 常量反向核它，而不是靠人读。
# 冻结态包内没有仓库根那份 README.md（产物里那份是打包时拷进去的），逐条显式 SKIP
# （源文件断言的老坑），"产物里到底有没有带 README.md"由闸门按实际产物目录核。
W_DOC_LABELS = ("W29 CLI 可照抄行里的长选项都在 argparse 里真实存在",
                "W29b argparse 里每个长选项都被文档覆盖到",
                "W29c 文档不是空壳：够长且八个必备小节齐",
                "W30 退出码表与 cli 常量一一对应",
                "W31 文档里不含真实口令或内网地址",
                "W31b 写清了 --prefix 的分隔符要自己写上",
                "W31c 写清了缺 UCRT 的症状、KB 号与包内修复脚本",
                "W31d 写清了另外两条线与各自系统门槛")
doc_path = os.path.join(ROOT, "README.md")
if not os.path.isfile(doc_path):
    for _lbl in W_DOC_LABELS:
        skip(_lbl, "冻结态包内没有仓库根的 README.md，只在开发态核")
else:
    doc = open(doc_path, encoding="utf-8").read()
    _long_opts = sorted({o for a in cli_mod.build_parser()._actions
                         for o in a.option_strings if o.startswith("--")})
    # 幽灵开关只核"能照抄的那几行"（含 excel2mysql-cli 的行）：文档别处合法地写着
    # build_exe.py 的 --fast、--keep-cache 和用例的 --no-parity，整篇扫会把正当提法判成错
    _cli_lines = [ln for ln in doc.split("\n") if "excel2mysql-cli" in ln]
    _ghost = sorted({t for ln in _cli_lines
                     for t in re.findall(r"--[a-z][a-z0-9-]*", ln) if t not in _long_opts})
    check(W_DOC_LABELS[0], not _ghost, _ghost or "无凭空写出的开关")
    _undoc = [o for o in _long_opts if o not in doc]
    check(W_DOC_LABELS[1], not _undoc, _undoc or "全部覆盖")
    _secs = ("前置条件", "怎么用最省事的四条", "支持哪些输入", "连接信息",
             "动作与选项", "日志形态与退出码", "最小验收", "跑不通时")
    # 长度下限按"开发向文档"取：README 只承载选型、环境、打包、启动、命令参考与闸门，
    # 界面操作细则在 help_doc.py（程序内《使用说明》），所以这里核的是命令参考没被砍空，
    # 不是整篇手册的字数
    check(W_DOC_LABELS[2], len(doc) > 9000 and all(s in doc for s in _secs),
          (len(doc), [s for s in _secs if s not in doc]))
    # 退出码只从「日志形态与退出码」那张表取：整篇扫会把别的表格首列数字当成退出码
    _sec = re.search(r"^### \d+\.5 日志形态与退出码.*?(?=^### |\Z)", doc, re.M | re.S)
    _codes_doc = sorted({int(x) for x in
                         re.findall(r"^\|\s*(\d{1,3})\s*\|", _sec.group(0), re.M)}) if _sec else []
    _codes_real = sorted({cli_mod.EXIT_OK, cli_mod.EXIT_USAGE, cli_mod.EXIT_PARSE,
                          cli_mod.EXIT_CONNECT, cli_mod.EXIT_IMPORT, cli_mod.EXIT_SCRIPT, 130})
    check(W_DOC_LABELS[3], bool(_sec) and _codes_doc == _codes_real,
          "文档里 %s / 代码里 %s" % (_codes_doc, _codes_real))
    # 只在一行内判断，别跨行：`password =` 换行后紧跟 `database =` 时，
    # 带 \s* 的正则会吃掉换行把下一个键名当成"口令值"误报。
    # 现场本来就要改这一行，所以"改成你的密码"这类占位不算泄漏，算的是像真凭据的形态
    _pw = [v.strip() for v in re.findall(r"^[ \t]*password[ \t]*=[ \t]*(\S.*)$", doc, re.M)]
    _leak = [v for v in _pw if not any(k in v for k in ("改成", "你的", "口令", "<", "留空"))]
    check(W_DOC_LABELS[4], not _leak and "Root@" not in doc and "172.18." not in doc, _leak)
    check(W_DOC_LABELS[5], "自己写上下划线" in doc, doc.count("--prefix"))
    check(W_DOC_LABELS[6],
          "KB2999226" in doc and "api-ms-win-crt" in doc and "只拷 exe 跑不起来" in doc
          and "win7_sp1_patch" in doc and "3.check_win7.bat" in doc
          and "4.install_minimal.bat" in doc,
          "文档要点到 win7_sp1_patch 里那几支脚本，现场才知道修复工具就在包里")
    check(W_DOC_LABELS[7],
          "Java 版" in doc and "Windows 10 1809" in doc and "三条线跑出来的表结构" in doc)


# ---------------------------------------------------------------------------
# --- W32-W36：建不出表的输入必须卡在解析阶段（生产实测抓出来的两类） ---
# W33 是这一类的反例，特意放在一起：只有表头的采集模板**不该**被拒（字段是设计好的表结构，
# 一行数据都还没填不是错误），它必须正常出表定义、并在日志里说清"整列空、建出来是空表"
from openpyxl import Workbook as _WB

_empty = os.path.join(_tmpdir, "空表（w_empty）.xlsx")
_wb = _WB()
_wb.active.title = "数据"
_wb.save(_empty)
_header_only = os.path.join(_tmpdir, "只有表头（w_header）.xlsx")
_wb = _WB()
_ws = _wb.active
_ws.title = "数据"
_ws.append(["编号", "名称"])
_wb.save(_header_only)
_fake = os.path.join(_tmpdir, "假excel（w_fake）.xlsx")
with open(_fake, "wb") as _fh:
    _fh.write(b"this is not a workbook")
_good = os.path.join(_tmpdir, "正常表（w_good）.xlsx")
_wb = _WB()
_ws = _wb.active
_ws.title = "数据"
_ws.append(["序号", "数值"])
_ws.append([1, 2.5])
_wb.save(_good)
# 表头带括号注释：真实表单里"发布日期（yyyy-MM-dd）""金额（元）"这类写法极常见，
# 注释是给填表人看的，不该进字段名（用户 2026-10-20 点名）
_annot = os.path.join(_tmpdir, "表头带注释（w_annot）.xlsx")
_wb = _WB()
_ws = _wb.active
_ws.title = "数据"
_ws.append([u"发布日期（yyyy-MM-dd）", u"发布日期", u"金额（元）"])
_ws.append([u"2008-05-19", u"2008-05-20", u"12.50"])
_wb.save(_annot)
# 表头整行都没写字：pandas 会连列宽一起丢掉，但"这张表有几列"文件里还记着
_blank_head = os.path.join(_tmpdir, "空表头模板（w_blankhead）.xlsx")
_wb = _WB()
_ws = _wb.active
_ws.title = "数据"
_ws.append(["", "", ""])
_wb.save(_blank_head)
# 口径串没写括号的写法：现场表头里 `发布日期：yyyy-MM-dd` 和 `发布日期 yyyy-MM-dd` 都很常见
_mask = os.path.join(_tmpdir, "表头挂日期格式（w_mask）.xlsx")
_wb = _WB()
_ws = _wb.active
_ws.title = "数据"
_ws.append([u"发布日期：yyyy-MM-dd", u"填表日期 yyyymmdd", u"发布日期(YYYYMMDD",
            u"累年月日降水量≥1.0mm平均日数"])
_ws.append([u"2008-05-19", u"20080520", u"20080521", u"5"])
_wb.save(_mask)

code, out, _con = run_cli(["--dry-run", _empty, "-d", "db1"])
check("W32 一张格子都没有的表被判解析失败（不给库里发非法建表语句）",
      code == cli_mod.EXIT_PARSE and "没有解析到任何字段" in out, out.strip()[:90])
code, out, _con = run_cli(["--dry-run", _header_only, "-d", "db1"])
check("W33 只有表头的采集模板不再被拒：字段 2 列、数据 0 行、日志点名空列与空表",
      code == cli_mod.EXIT_OK and "没有解析到任何字段" not in out
      and "字段 2 列 | 数据 0 行" in out
      and "整列为空的字段照旧保留：编号、名称" in out
      and "没有数据行，建出来是空表" in out, out.strip()[-300:])
code, out, _con = run_cli(["--dry-run", _annot, "-d", "db1"])
check("W33b 表头括号注释不进字段名，剥完同名靠 _2 去重，注释里仍留原文",
      code == cli_mod.EXIT_OK
      and "`fa_bu_ri_qi`" in out and "`fa_bu_ri_qi_2`" in out and "`jin_e`" in out
      and not re.search(r"`[a-z0-9_]*yyyy", out)
      and "COMMENT '发布日期（yyyy-MM-dd）'" in out, out.strip()[-260:])
code, out, _con = run_cli(["--dry-run", _blank_head, "-d", "db1"])
check("W33c 表头整行没写字也出 DDL：按列宽兜底成 col、col_2、col_3，并在日志点名",
      code == cli_mod.EXIT_OK
      and "`col`" in out and "`col_2`" in out and "`col_3`" in out
      and "没有解析到任何字段" not in out
      and "字段 3 列 | 数据 0 行" in out
      and "表头整行没写字，按列宽兜底成：col、col_2、col_3" in out, out.strip()[-300:])
code, out, _con = run_cli(["--dry-run", _mask, "-d", "db1"])
check("W33d 没括号的日期格式串也不进字段名，且不误吃单位字母 mm",
      code == cli_mod.EXIT_OK
      and "`fa_bu_ri_qi`" in out and "`tian_biao_ri_qi`" in out
      and "`lnyrjsl10mmpjrs`" in out
      and not re.search(r"`[a-z0-9_]*(yyyy|mmdd|mm_dd)", out), out.strip()[-300:])
code, out, _con = run_cli(["--dry-run", _fake, "-d", "db1"])
check("W34 扩展名对但内容不是 Excel：拦在解析阶段而不是 SQL 语法错误",
      code == cli_mod.EXIT_PARSE and "解析失败" in out, out.strip()[:90])
code, out, _con = run_cli(["--dry-run", _good, _fake, "-d", "db1"])
check("W35 一批里混一张坏表：好表照常出 DDL，坏表单独记账并抬退出码",
      code == cli_mod.EXIT_PARSE and "w_good" in out and "解析失败" in out,
      out.strip()[-120:])
code, out, _con = run_cli([_fake, "-d", "db1"])
check("W36 真导入路径同样不透出 CREATE TABLE（坏表不会去碰库）",
      code in (cli_mod.EXIT_PARSE, cli_mod.EXIT_IMPORT)
      and "CREATE TABLE" not in out, code)

# W37：真实数据包里的 Excel 多是生成器写的"最小样式表"，openpyxl 读它会 warn
# "Workbook contains no default style"，每张表一行混进日志。那既不是本工具的日志格式，
# 也不影响结果，只会把真正的错误淹掉——日志里不该出现第三方警告。
import zipfile as _zip

_style_raw = os.path.join(_tmpdir, "_style_raw.xlsx")
_wb = _WB()
_ws = _wb.active
_ws.title = "数据"
_ws.append(["编号", "名称"])
_ws.append([1, "甲"])
_ws.append([2, "乙"])
_wb.save(_style_raw)
_no_style = os.path.join(_tmpdir, "无默认样式（w_style）.xlsx")
_zin = _zip.ZipFile(_style_raw, "r")
_zout = _zip.ZipFile(_no_style, "w", _zip.ZIP_DEFLATED)
for _item in _zin.infolist():
    _data = _zin.read(_item.filename)
    if _item.filename == "xl/styles.xml":
        _text = _data.decode("utf-8")
        _s = _text.find("<cellStyles")
        if _s >= 0:
            _e = _text.find("</cellStyles>", _s)
            _text = _text[:_s] + _text[_e + len("</cellStyles>"):]
        _data = _text.encode("utf-8")
    _zout.writestr(_item, _data)
_zout.close()
_zin.close()
os.remove(_style_raw)
import warnings as _warn

import openpyxl as _oxl

with _warn.catch_warnings(record=True) as _caught:
    _warn.simplefilter("always")
    _oxl.load_workbook(_no_style)
check("W37 夹具确实是没有命名样式的那种表（openpyxl 会 warn）",
      any("no default style" in str(_w.message) for _w in _caught),
      [str(_w.message)[:40] for _w in _caught])
code, out, _con = run_cli(["--dry-run", _no_style, "-d", "db1"])
check("W37b 这种表照样出 DDL（压警告没把读表一起压坏）",
      code == cli_mod.EXIT_OK and "CREATE TABLE" in out, out.strip()[:90])

# W37c 关键的一条：第三方警告不能出现在 stderr 上。
# run_cli 只看 stdout 与 --log 那个日志文件，而警告是打到 stderr 的——
# 断"日志文件里没有警告行"是条永真断言，把 cli.py 的过滤器全删了它照样绿。
# 所以这里自己抓 stderr，并且先清掉 openpyxl 的 __warningregistry__：
# 默认动作是 once，同一处只打一次，不清就测不到第二次。
import contextlib as _ctx
import openpyxl.styles.stylesheet as _ss

_LOG = os.path.join(_tmpdir, "logs", "w37c.log")
_seq = os.path.dirname(_LOG)
os.makedirs(_seq, exist_ok=True)
if os.path.isfile(_LOG):
    os.remove(_LOG)
_err = io.StringIO()
_reg = getattr(_ss, "__warningregistry__", None)
if isinstance(_reg, dict):
    _reg.clear()
with _ctx.redirect_stderr(_err):
    with _ctx.redirect_stdout(io.StringIO()):
        cli_mod.run(["--dry-run", _no_style, "-d", "db1", "--log", _LOG])
check("W37c 读这种表的 stderr 上没有第三方警告行",
      "no default style" not in _err.getvalue() and "UserWarning" not in _err.getvalue(),
      _err.getvalue().strip()[:120])
# 同一段 stderr 里，与本工具无关但属于别类的警告还是要打出来——压的是那两类，不是全部
_probe = io.StringIO()
with _ctx.redirect_stderr(_probe):
    _warn.warn("别的警告照常打出来", UserWarning)
check("W37d 只压那两类第三方警告，没有把警告整体静音",
      "别的警告照常打出来" in _probe.getvalue(), _probe.getvalue().strip()[:80])

# 命名开关与"空列按字段名定类型"（用户 2026-10-20 第 4 件；默认行为必须一个字都不变）
_naming = os.path.join(_tmpdir, "水库信息表.xlsx")     # 不带括号：表名才走中文转写
_wb = _WB()
_ws = _wb.active
_ws.title = "数据"
_ws.append([u"订单金额", u"备注说明"])
_ws.append([12.5, u"甲"])
_wb.save(_naming)
code, out, _con = run_cli(["--dry-run", _naming, "-d", "db1"])
check("W38 默认全拼：不开开关时表名与字段名一个字都不变",
      code == cli_mod.EXIT_OK and "`shui_ku_xin_xi_biao`" in out
      and "`ding_dan_jin_e`" in out, out.strip()[-160:])
code, out, _con = run_cli(["--initials", "--dry-run", _naming, "-d", "db1"])
check("W39 --initials：表名与字段名取首字母，启动行写明命名规则",
      code == cli_mod.EXIT_OK and "`skxxb`" in out and "`ddje`" in out
      and "命名取拼音首字母" in out, out.strip()[:130])
_types = os.path.join(_tmpdir, "采集模板（w_types）.xlsx")
_wb = _WB()
_ws = _wb.active
_ws.title = "数据"
_ws.append([u"发布日期", u"更新时间", u"订单金额", u"数量", u"备注说明"])
_wb.save(_types)
code, out, _con = run_cli(["--dry-run", _types, "-d", "db1"])
check("W40 空模板按字段名定类型：日期/时间/金额/数量各归各位且留痕",
      code == cli_mod.EXIT_OK
      and "`fa_bu_ri_qi` DATE" in out and "`geng_xin_shi_jian` DATETIME" in out
      and "`ding_dan_jin_e` DECIMAL(16,2)" in out and "`shu_liang` INT" in out
      and "该列没有值" in out, out.strip()[-320:])
code, out, _con = run_cli(["--upper", "--dry-run", _naming, "-d", "db1"])
check("W41 --upper：表名与字段名转大写，启动行写明大小写规则；默认小写留原样",
      code == cli_mod.EXIT_OK and "`SHUI_KU_XIN_XI_BIAO`" in out
      and "`DING_DAN_JIN_E`" in out and "表名与字段名转大写" in out
      and "`shui_ku_xin_xi_biao`" not in out, out.strip()[:180])
code, out, _con = run_cli(["--initials", "--upper", "--dry-run", _naming, "-d", "db1"])
check("W41b --initials 与 --upper 叠加：先取首字母再转大写",
      code == cli_mod.EXIT_OK and "`SKXXB`" in out and "`DDJE`" in out
      and "命名取拼音首字母" in out and "表名与字段名转大写" in out,
      out.strip()[:150])
code, out, _con = run_cli(["--upper", "--dry-run", _types, "-d", "db1"])
check("W41c 大写时空列按名定类型照旧，字段名整体大写",
      code == cli_mod.EXIT_OK and "`FA_BU_RI_QI` DATE" in out
      and "`GENG_XIN_SHI_JIAN` DATETIME" in out
      and "`DING_DAN_JIN_E` DECIMAL(16,2)" in out
      and "`SHU_LIANG` INT" in out, out.strip()[-320:])
_kinds = os.path.join(_tmpdir, "采集模板（w_kinds）.xlsx")
_wb = _WB()
_ws = _wb.active
_ws.title = "数据"
_ws.append([u"年龄", u"系数", u"评分", u"经度", u"长度", u"是否有效", u"序号", u"年份"])
_wb.save(_kinds)
code, out, _con = run_cli(["--dry-run", _kinds, "-d", "db1"])
check("W42 空模板中文词表扩展：年龄/系数/评分/经纬度/计量/布尔/序号/年份各归各位",
      code == cli_mod.EXIT_OK
      and "`nian_ling` INT" in out and "`xi_shu` DECIMAL(18,6)" in out
      and "`ping_fen` DECIMAL(10,2)" in out and "`jing_du` DECIMAL(18,8)" in out
      and "`chang_du` DECIMAL(18,4)" in out
      and "`shi_fou_you_xiao` TINYINT(1)" in out
      and "`xu_hao` INT" in out and "`nian_fen` INT" in out
      and "含比率/系数语义" in out, out.strip()[-480:])

section("X CLI 端到端·假连接（断到真实落库调用序列，不只看退出码）")


class FakeCursor:
    def __init__(self, rec, fail_on=None, fail_if=None):
        self.rec = rec
        self.fail_on = fail_on
        self.fail_if = fail_if
        self.rowcount = -1

    def _should_fail(self, sql):
        if self.fail_if is not None:
            return bool(self.fail_if(sql))
        return bool(self.fail_on and self.fail_on in sql)

    def execute(self, sql, params=None):
        self.rec.setdefault("execute", []).append((sql, params))
        # 模拟驱动：SELECT/SHOW 才有 description，CLI 靠它判断要不要 fetchall
        self.description = ["id"] if sql.strip().upper().startswith(("SELECT", "SHOW")) else None
        if self._should_fail(sql):
            raise RuntimeError("fake boom on %s" % sql[:30])

    def close(self):
        self.rec["cursor_closed"] = True

    def executemany(self, sql, seq):
        # 插入失败必须在这里也能抛：core.insert_data 走的是 executemany，
        # 只在 execute 上设失败点会连"导入失败回滚"这条路径都没测到。
        if self._should_fail(sql):
            raise RuntimeError("fake boom on %s" % sql[:30])
        self.rec.setdefault("executemany", []).append((sql, list(seq)))
        self.rowcount = len(list(seq))

    def fetchall(self):
        self.rec["fetchall"] = self.rec.get("fetchall", 0) + 1
        return [("testdb",), ("information_schema",)]

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


class FakeConn:
    def __init__(self, rec, fail_ddl=False, fail_insert=False, fail_if=None):
        self.rec = rec
        self.fail_ddl = fail_ddl
        self.fail_insert = fail_insert
        self.fail_if = fail_if

    def cursor(self):
        fail_on = "CREATE TABLE" if self.fail_ddl else ("INSERT INTO" if self.fail_insert else None)
        return FakeCursor(self.rec, fail_on=fail_on, fail_if=self.fail_if)

    def begin(self):
        self.rec.setdefault("begin", 0)
        self.rec["begin"] += 1

    def commit(self):
        self.rec.setdefault("commit", 0)
        self.rec["commit"] += 1

    def rollback(self):
        self.rec.setdefault("rollback", 0)
        self.rec["rollback"] += 1

    def close(self):
        self.rec.setdefault("close", 0)
        self.rec["close"] += 1


def with_fake_conn(rec, fail_ddl=False, fail_insert=False, cfg_seen=None, fail_ddl_for=None,
                   fail_if=None):
    """把 cli 的 connect 换成假连接：记录调用序列，同时抓住连接配置。

    fail_ddl_for 给一个表名时只让那张表的 CREATE 失败，用来验"单表失败不阻断其余表"；
    fail_if 给一个函数时按语句内容决定失败点（SQL 脚本线要的是"第 N 条语句被拒"）。
    """
    if fail_if is None and fail_ddl_for:
        fail_if = lambda sql: sql.startswith("CREATE TABLE `%s`" % fail_ddl_for)

    def _fake(cfg):
        if cfg_seen is not None:
            cfg_seen.update(cfg)
        return FakeConn(rec, fail_ddl=fail_ddl, fail_insert=fail_insert, fail_if=fail_if)
    return _fake


def cli_fake(argv, fail_if=None, fail_ddl_for=None):
    """通用假连接跑法：跑 cli.run，返回 (code, 日志, 调用记录, 连接配置)。"""
    rec = {}
    seen = {}
    real_connect = cli_mod.connect
    cli_mod.connect = with_fake_conn(rec, cfg_seen=seen, fail_if=fail_if,
                                     fail_ddl_for=fail_ddl_for)
    try:
        code, out, _ = run_cli(argv)
    finally:
        cli_mod.connect = real_connect
    return code, out, rec, seen


def cli_import(fixture, extra=None, fail_ddl=False, fail_insert=False, fail_ddl_for=None):
    """跑一次真实导入流程但连接是假的，返回 (code, 日志, 调用记录, 连接配置)。"""
    rec = {}
    seen = {}
    real_connect = cli_mod.connect
    cli_mod.connect = with_fake_conn(rec, fail_ddl, fail_insert, seen, fail_ddl_for)
    try:
        code, out, _ = run_cli([fixture, "-d", "demo"] + (extra or []))
    finally:
        cli_mod.connect = real_connect
    return code, out, rec, seen


code, out, rec, seen = cli_import(v_orders)
check("X1 单文件导入退出码 0", code == cli_mod.EXIT_OK, "%s\n%s" % (code, out[-400:]))
check("X2 连接配置传到 core.connect",
      (seen.get("host"), seen.get("database")) == ("127.0.0.1", "demo"), seen)
check("X3 每张表先 begin 再建表", rec.get("begin") == 1, rec.get("begin"))
_ddl = [s for s, _ in rec.get("execute", []) if s.startswith("CREATE TABLE")]
check("X4 建表 DDL 实际发出去了", len(_ddl) == 1 and "`ods_t_order`" in _ddl[0], _ddl)
check("X5 建表与插入各自提交（core.create_table / insert_data 都 commit，共两次）",
      rec.get("commit") == 2, rec.get("commit"))
_em = rec.get("executemany", [])
check("X6 插入走 executemany 且批大小 500", len(_em) == 1 and len(_em[0][1]) == 3,
      [(e[0][:40], len(e[1])) for e in _em])
check("X7 插入参数按列序绑定（不是列名错配）",
      _em and _em[0][1][0][0] == "SO0001" and _em[0][1][0][2] == 12.5,
      _em[0][1][0] if _em else None)
check("X8 日期串已转成 date 对象再交给驱动",
      hasattr(_em[0][1][0][3], "year") if _em else False, _em[0][1][0][3] if _em else None)
check("X9 连接被关闭（不泄漏句柄）", rec.get("close") == 1, rec.get("close"))
check("X10 日志里有逐表导入行数", "导入完成 | `ods_t_order` 共 3 行" in out, out.strip()[-200:])
check("X11 汇总行给出成功表数与总行数",
      "导入汇总 | 成功 1 张表 / 3 行" in out, out.strip()[-160:])

code, out, rec, _ = cli_import(v_orders, extra=["--dry-run"])
check("X12 dry-run 完全不建连接",
      code == cli_mod.EXIT_OK and not rec and "CREATE TABLE" in out, (code, rec))
check("X13 dry-run 明说没有连接 MySQL", "本次没有连接 MySQL" in out, out.strip()[-120:])
check("X14 dry-run 不落库（没有 executemany）", "executemany" not in rec, rec.keys())

code, out, rec, _ = cli_import(v_orders, extra=["--drop-if-exists"])
_ddl2 = [s for s, _ in rec.get("execute", []) if "DROP TABLE" in s or "CREATE TABLE" in s]
check("X15 --drop-if-exists 把 DROP 拼进同一次建表调用",
      any(s.startswith("DROP TABLE IF EXISTS `ods_t_order`") for s, _ in rec.get("execute", [])),
      _ddl2)
check("X16 --drop-if-exists 显式警告会删同名表",
      "危险开关" in out and "同名表会先被删除" in out, out.strip()[:120])

code, out, rec, _ = cli_import(v_orders, fail_ddl=True)
check("X17 建表被拒时退出码 4", code == cli_mod.EXIT_IMPORT, code)
check("X18 建表失败后没有执行插入", "executemany" not in rec, rec.keys())
check("X19 建表失败把原文 DDL 打进日志", "失败的建表语句" in out and "CREATE TABLE `ods_t_order`" in out)
check("X20 建表失败提示 DDL 不随事务回滚", "隐式提交" in out, out.strip()[:150])

code, out, rec, _ = cli_import(v_orders, fail_insert=True)
check("X21 插入失败退出码 4", code == cli_mod.EXIT_IMPORT, code)
check("X22 插入失败触发 rollback（本表数据不半写）", rec.get("rollback") == 1, rec)
check("X23 插入失败日志写明回滚与行数",
      "行数据未写入，本表事务已回滚" in out and "3 行" in out, out.strip()[:160])

conn_err = {}


def _boom(cfg):
    conn_err.update(cfg)
    import pymysql
    raise pymysql.err.OperationalError(1049, "Unknown database 'nosuch'")


_real = cli_mod.connect
cli_mod.connect = _boom
try:
    code, out, _ = run_cli([v_orders, "-d", "nosuch"])
finally:
    cli_mod.connect = _real
check("X24 连接失败退出码 3（区别于导入失败 4）", code == cli_mod.EXIT_CONNECT, code)
check("X25 连接失败展开 MySQL 错误码 1049",
      "MySQL 错误码 1049" in out and "Unknown database" in out, out.strip()[:160])
check("X26 连接失败带上当时的目标串（user@host:port/db）",
      "root@127.0.0.1:3306/nosuch" in out, out.strip()[:120])

batch_dir = os.path.join(_tmpdir, "batch")
shutil.rmtree(batch_dir, ignore_errors=True)
os.makedirs(batch_dir)
make_xlsx("库存.xlsx", pd.DataFrame({"物料编码": ["A1", "A2"], "库存量": [10, 20]}),
          folder=batch_dir)
make_xlsx("人员名册.xlsx", pd.DataFrame({"工号": ["E1"], "姓名": ["甲"]}), folder=batch_dir)
code, out, rec, _ = cli_import(batch_dir)
check("X27 目录批量导入：两张表各自 begin，连接只关一次",
      code == cli_mod.EXIT_OK and rec.get("begin") == 2 and rec.get("close") == 1,
      (code, rec.get("begin"), rec.get("close")))
check("X27b 日志逐表报序号 [1/2] [2/2]（现场能判断卡在哪张）",
      "开始处理 | [1/2]" in out and "开始处理 | [2/2]" in out, out[:60])
check("X27c 汇总行给出 2 张表 / 3 行", "导入汇总 | 成功 2 张表 / 3 行" in out,
      out.strip()[-120:])

code, out, rec, _ = cli_import(batch_dir, fail_ddl_for="ku_cun")
_create_sqls = [s for s, _ in rec.get("execute", []) if s.startswith("CREATE TABLE")]
check("X28 单张表建表被拒不阻断其余表（另一张照常建+导入）",
      code == cli_mod.EXIT_IMPORT and rec.get("begin") == 2 and len(_create_sqls) == 2
      and "导入完成" in out, (code, rec.get("begin"), len(_create_sqls)))
check("X28b 失败的那张在汇总里被点名", "失败 1 张" in out and "ku_cun" in out,
      out.strip()[-160:])

code, out, rec, _ = cli_import(v_orders, extra=["--prefix", "ods_"])
check("X29 --prefix 保留末尾下划线作分隔符（曾经的真 bug）",
      any("`ods_ods_t_order`" in s for s, _ in rec.get("execute", [])),
      [s[:60] for s, _ in rec.get("execute", [])])
code, out, rec, _ = cli_import(v_orders, extra=["--prefix", "ods"])
check("X30 --prefix 不带下划线时告警而不是默默连写",
      "分隔符" in out and "直接接在原名前面" in out, out.strip()[:140])
code, out, rec, _ = cli_import(v_orders, extra=["--prefix", ""])
check("X31 --prefix 留空等于移除前缀并提示",
      "前缀为空" in out and any("`ods_t_order`" in s for s, _ in rec.get("execute", [])), out[:120])

swap_src = make_xlsx("库存（ods_t_stock）.xlsx", pd.DataFrame({"物料编码": ["A1"]}))
code, out, rec, _ = cli_import(swap_src, extra=["--swap-prefix", "dim_"])
check("X32 --swap-prefix 换掉首段：ods_t_stock → dim_t_stock",
      any("CREATE TABLE `dim_t_stock`" in s for s, _ in rec.get("execute", [])),
      [s[:60] for s, _ in rec.get("execute", [])])
check("X33 换首段在日志里报告改了哪张", "首段前缀替换" in out and "ods_t_stock" in out, out[:160])

nosep = make_xlsx("无下划线（users）.xlsx", pd.DataFrame({"编号": [1]}))
code, out, rec, _ = cli_import(nosep, extra=["--swap-prefix", "dim"])
check("X34 名里没有下划线的表不改并逐条说明",
      "跳过 1 张" in out, out[out.find("首段") if "首段" in out else 0:][:200])

two_same = os.path.join(_tmpdir, "两同名")
shutil.rmtree(two_same, ignore_errors=True)
os.makedirs(two_same, exist_ok=True)
shutil.copy(v_orders, os.path.join(two_same, "a-订单（ods_t_order）.xlsx"))
shutil.copy(v_orders, os.path.join(two_same, "b-订单（ods_t_order）.xlsx"))
code, out, rec, _ = cli_import(two_same)
_x35 = [s.split(" (")[0] for s, _ in rec.get("execute", []) if s.startswith("CREATE TABLE")]
check("X35 同批撞名自动 _2（断真实发出去的表名，不是另算一遍）",
      _x35 == ["CREATE TABLE `ods_t_order`", "CREATE TABLE `ods_t_order_2`"], _x35)
check("X36 撞名两张表都真的建了（begin 两次、退出码 0）",
      rec.get("begin") == 2 and code == cli_mod.EXIT_OK, (rec.get("begin"), code))

code, out, _ = run_cli([os.path.join(_tmpdir, "没有这个文件.xlsx"), "-d", "demo"])
check("X37 路径不存在判用法错误并点名", code == cli_mod.EXIT_USAGE, (code, out[:200]))
check("X38 路径不存在时日志里点名「路径不存在」", "路径不存在" in out, out[:200])

bad_ext = os.path.join(_tmpdir, "readme.txt")
with open(bad_ext, "w", encoding="utf-8") as fh:
    fh.write("not excel")
code, out, _ = run_cli([bad_ext, "-d", "demo"])
check("X39 直接指定非 Excel 文件：不支持类型 → 退出码 1",
      code == cli_mod.EXIT_USAGE and "不支持的类型" in out, (code, out[:120]))


# ---------------------------------------------------------------------------
section("Y CLI 输入收集、压缩包与日志形态")

demo_dir = os.path.join(_tmpdir, "demo_tree")
shutil.rmtree(demo_dir, ignore_errors=True)
os.makedirs(os.path.join(demo_dir, "子目录"))
make_xlsx("一层.xlsx", pd.DataFrame({"编号": [1]}), folder=demo_dir)
make_xlsx("二层.xlsx", pd.DataFrame({"编号": [2]}), folder=os.path.join(demo_dir, "子目录"))
with open(os.path.join(demo_dir, "忽略.csv"), "w", encoding="utf-8") as fh:
    fh.write("a,b\n")

entries, zips, _sqls = cli_mod.collect_inputs([demo_dir], True, cli_mod.Log(quiet=True))
check("Y1 递归目录收齐 xlsx", len(entries) == 2, [e[1] for e in entries])
check("Y2 目录里的 zip 单独归类", zips == [], zips)
check("Y3 非 Excel 文件不进待处理列表",
      all(not e[0].endswith(".csv") for e in entries), [e[0] for e in entries])
entries2, _, _ = cli_mod.collect_inputs([demo_dir], False, cli_mod.Log(quiet=True))
check("Y4 --no-recursive 只收一层", len(entries2) == 1, [e[1] for e in entries2])

zip_tree = os.path.join(_tmpdir, "ziptree")
os.makedirs(zip_tree, exist_ok=True)
zpath = os.path.join(zip_tree, "批量.zip")
inner = make_xlsx("包内一层.xlsx", pd.DataFrame({"编号": [3]}), folder=_tmpdir)
with zipfile.ZipFile(zpath, "w") as zf:
    zf.write(inner, "深层/包内一层.xlsx")
code, out, rec, _ = cli_import(zpath)
check("Y5 压缩包直接作为入参：展开后成功导入", code == cli_mod.EXIT_OK, (code, out[-300:]))
check("Y6 日志把包内路径写成 包名!包内/文件", "批量.zip!深层/包内一层.xlsx" in out, out[:260])

entries3, zips3, _sql3 = cli_mod.collect_inputs([zip_tree], True, cli_mod.Log(quiet=True))
check("Y7 目录里的 zip 被收进 zips 而不是当 Excel",
      len(zips3) == 1 and not entries3, (entries3, zips3))
check("Y8 不存在的输入路径被显式报错（不静默跳过）", True, "见 X37")

code, out, rec, _ = cli_import(merged_path)
check("Y9 合并单元格文件不静默丢：退出码 2 且日志给原因",
      code == cli_mod.EXIT_PARSE and "合并单元格" in out, (code, out[:200]))
check("Y10 合并单元格拒绝后没有建任何表", "execute" not in rec, rec.keys())

code, out, _ = run_cli([merged_two, "-d", "demo", "--dry-run"])
check("Y10a 两行合并表头在命令行照常导入，并说明按几行展开",
      code == cli_mod.EXIT_OK and "合并表头按 2 行展开" in out, (code, out[:400]))

# 日志形态：一行自带判读，行首符号 + 大写标签，分段 | ；错误档必须能单独 grep
code, out, _ = run_cli([v_orders, "-d", "demo", "--dry-run"])
lines = log_lines(out)
_records = [ln for ln in lines if not ln.startswith("      ")]
check("Y10b 每条记录行都带时间戳（缩进的详情行除外，事后能对上发生时刻）",
      bool(_records) and all(
          re.match(r"^\[\d{4}-\d\d-\d\d \d\d:\d\d:\d\d\] ", ln)
          for ln in out.splitlines() if ln.strip() and not ln.startswith("      ")),
      out.splitlines()[:1])
check("Y11 每条记录行都以「符号 [大写标签]」开头（详情行另计）",
      all(re.match(r"^[+x!-] \[[A-Z_]+\] ", ln) for ln in _records),
      [ln for ln in _records if not re.match(r"^[+x!-] \[[A-Z_]+\] ", ln)][:2])
check("Y12 成功与错误分档不同符号",
      any(ln.startswith("+ [OK]") for ln in lines) and "x [ERROR]" not in "\n".join(lines))
check("Y13 详情行用缩进跟在记录后，不冒充新记录",
      any(ln.startswith("      ") for ln in lines))
code, out, _ = run_cli([merged_path, "-d", "demo"])
check("Y14 错误可单独 grep：x [ERROR] 一档捞得出来",
      sum(1 for ln in log_lines(out) if ln.startswith("x [ERROR]")) >= 1, out[:200])

qcode, qout, qconsole = run_cli([v_orders, "-d", "demo", "--dry-run", "-q"])
check("Y15 --quiet 压掉控制台的记录行，日志文件仍记全量",
      qcode == cli_mod.EXIT_OK
      and not [ln for ln in qconsole.splitlines() if re.match(r"^[+x!-] \[", ln)]
      and any(ln.startswith("+ [OK]") for ln in log_lines(qout)),
      (qcode, qconsole.strip()[:60]))

qcode, qout, qconsole = run_cli([merged_path, "-d", "demo", "-q"])
check("Y15b --quiet 不会把错误连同它的详情一起藏掉",
      qcode == cli_mod.EXIT_PARSE
      and any(ln.startswith("x [ERROR]") for ln in qconsole.splitlines())
      and any(ln.startswith("      原因：") for ln in qconsole.splitlines()),
      qconsole.strip()[:120])

with clean_env():
    lg = cli_mod.Log(path=os.path.join(_tmpdir, "append.log"))
    lg.ok("第一次")
    lg.close()
    lg2 = cli_mod.Log(path=os.path.join(_tmpdir, "append.log"))
    lg2.error("第二次")
    lg2.close()
    with open(os.path.join(_tmpdir, "append.log"), encoding="utf-8") as fh:
        alllog = fh.read()
check("Y16 日志文件是追加不是覆盖（一次任务可能跑多轮）",
      "第一次" in alllog and "第二次" in alllog and alllog.count("\n") >= 2, alllog)
check("Y17 日志文件写 UTF-8（中文不乱码）", "建表" in alllog or "第一次" in alllog)

check("Y18 退出码常量互不相同（脚本可按码分流）",
      len({cli_mod.EXIT_OK, cli_mod.EXIT_USAGE, cli_mod.EXIT_PARSE,
           cli_mod.EXIT_CONNECT, cli_mod.EXIT_IMPORT}) == 5,
      [cli_mod.EXIT_OK, cli_mod.EXIT_USAGE, cli_mod.EXIT_PARSE, cli_mod.EXIT_CONNECT, cli_mod.EXIT_IMPORT])
# 混合输入：两好一坏（坏的是合并单元格那个）。退出码规则在这里钉死——
# 解析有失败但动作全成功 → 2；解析和动作都失败 → 4，让调用方按"最严重的一档"分流。
mix_dir = os.path.join(_tmpdir, "mix")
shutil.rmtree(mix_dir, ignore_errors=True)
os.makedirs(mix_dir)
make_xlsx("好1.xlsx", pd.DataFrame({"编号": [1]}), folder=mix_dir)
make_xlsx("好2.xlsx", pd.DataFrame({"编号": [2]}), folder=mix_dir)
shutil.copy(merged_path, os.path.join(mix_dir, "坏.xlsx"))
code, out, rec, _ = cli_import(mix_dir)
check("Y19 解析有失败但导入全成功 → 退出码 2（不报 0 骗人）",
      code == cli_mod.EXIT_PARSE and rec.get("begin") == 2, (code, rec.get("begin")))
check("Y19b 日志同时给出失败项数和继续处理的表数",
      "解析阶段有失败" in out and "1 项未进入待导入列表" in out, out[:220])
code, out, rec, _ = cli_import(mix_dir, fail_ddl=True)
check("Y20 解析失败与建表失败同时发生 → 报更严重的 4（动作失败优先）",
      code == cli_mod.EXIT_IMPORT, (code, out[-160:]))

# Ctrl+C 要真的模拟一次中断：读源码里有没有 "130" 是个假断言，冻结态没有源码还会误红
_real_collect = cli_mod.collect_inputs


def _interrupt(*a, **k):
    raise KeyboardInterrupt()


cli_mod.collect_inputs = _interrupt
try:
    code, out, _ = run_cli([v_orders, "-d", "demo"])
finally:
    cli_mod.collect_inputs = _real_collect
check("Y21 跑到一半按 Ctrl+C 返回 130（与用法错误 1 区分开）", code == 130, (code, out[:120]))
check("Y21b 中止时日志明说剩余的表未执行", "已中止" in out, out[:160])

# --keep-temp 与 --batch 此前只写在 README 里、没有实测过，这两条把它们钉住
code, out, _ = run_cli([zpath, "--dry-run", "--keep-temp"])
kept = re.search(r"临时目录 \| 已保留解压出的 Excel：(.+)$", out.replace("\r", ""), re.M)
check("Y22 --keep-temp 报出临时目录，且解压出的 Excel 真留在那儿",
      code == cli_mod.EXIT_OK and bool(kept) and os.path.isdir(kept.group(1).strip())
      and any(f.endswith(".xlsx") for _, _, fs in os.walk(kept.group(1).strip()) for f in fs),
      kept.group(1).strip() if kept else out[-200:])
code, out, _ = run_cli([zpath, "--dry-run"])
gone = re.search(r"压缩包展开 \| .*?Excel 1 个", out)
check("Y22b 默认不留临时文件（跑完就清，不越堆越多）",
      code == cli_mod.EXIT_OK and gone and "已保留" not in out, out[-160:])

code, out, rec, _ = cli_import(v_orders, extra=["--batch", "2"])
_batches = [len(rows) for _, rows in rec.get("executemany", [])]
check("Y23 --batch 2 把 3 行真拆成 2 批（2+1），不是整表一次性发过去",
      _batches == [2, 1], _batches)
code, out, rec, _ = cli_import(v_orders)
check("Y23b 默认批大小 500：3 行就是一批",
      [len(rows) for _, rows in rec.get("executemany", [])] == [3], rec.get("executemany"))


# ---------------------------------------------------------------------------
section("Z SQL 脚本线（GUI「SQL 脚本改名」页的命令行版）")

sql_dir = os.path.join(_tmpdir, "sqls")
shutil.rmtree(sql_dir, ignore_errors=True)
os.makedirs(sql_dir)
sql_a = os.path.join(sql_dir, "建表.sql")
with open(sql_a, "w", encoding="utf-8") as fh:
    fh.write("-- 中文注释里有个 ; 分号，不该被当成语句边界\n"
             "CREATE TABLE `ods_t_order` (`id` BIGINT NOT NULL, PRIMARY KEY (`id`));\n"
             "CREATE TABLE `ods_t_user` (`id` BIGINT);\n"
             "INSERT INTO `ods_t_user` VALUES (1);\n"
             "INSERT INTO `ods_t_user` VALUES (2);\n")
sql_empty = os.path.join(sql_dir, "空脚本.sql")
with open(sql_empty, "w", encoding="utf-8") as fh:
    fh.write("-- 只有注释，没有语句\n")
sql_blank = os.path.join(_tmpdir, "真空白.sql")
with open(sql_blank, "w", encoding="utf-8") as fh:
    fh.write("")
sql_nohd = os.path.join(_tmpdir, "无下划线.sql")
with open(sql_nohd, "w", encoding="utf-8") as fh:
    fh.write("CREATE TABLE `lonename` (`id` INT);\n")
sql_gbk = os.path.join(sql_dir, "gbk脚本.sql")
with open(sql_gbk, "wb") as fh:
    fh.write("-- 记事本 ANSI 存的中文注释\nCREATE TABLE `gbk_one` (`id` INT);\n".encode("gbk"))
sql_sel = os.path.join(_tmpdir, "带查询.sql")
with open(sql_sel, "w", encoding="utf-8") as fh:
    fh.write("CREATE TABLE `sel_one` (`id` INT);\nSELECT `id` FROM `sel_one`;\n")
sql_bad = os.path.join(_tmpdir, "坏语句.sql")
with open(sql_bad, "w", encoding="utf-8") as fh:
    fh.write("CREATE TABLE `good_one` (`id` INT);\nTHIS IS NOT SQL;\n"
             "CREATE TABLE `never_runs` (`id` INT);\n")

ents, zps, sps = cli_mod.collect_inputs([sql_dir], True, cli_mod.Log(quiet=True))
check("Z1 .sql 归类到脚本，不混进 Excel 列表",
      len(sps) == 3 and not ents and not zps, (len(ents), len(sps)))

code, out, rec, seen = cli_fake([sql_a, "-d", "demo", "--run-sql"])
check("Z2 识别 CREATE TABLE 的数量与表名",
      "CREATE TABLE 2 处 / 表 2 张：ods_t_order(1)、ods_t_user(1)" in out, out[:150])
check("Z3 注释里的分号没被误切：正好 4 条语句",
      len(rec.get("execute", [])) == 4, [s[:30] for s, _ in rec.get("execute", [])])
check("Z4 每个文件一个事务：begin→执行→commit→close",
      rec.get("begin") == 1 and rec.get("commit") == 1 and rec.get("close") == 1, rec)
check("Z4b 连接配置照常传下去", seen.get("database") == "demo", seen)
check("Z4c 全部成功才提交，日志报出提交条数",
      "脚本提交 | 建表.sql | 4 条语句已提交" in out and "脚本汇总 | 成功 1 个文件，失败 0 个" in out,
      out[-200:])

code, out, rec, _ = cli_fake([sql_sel, "-d", "demo", "--run-sql"])
check("Z5 SELECT 语句会取空结果集（游标不挂着）",
      len(rec.get("execute", [])) == 2 and rec.get("fetchall") == 1,
      (len(rec.get("execute", [])), rec.get("fetchall")))

code, out, rec, _ = cli_fake([sql_a, "-d", "demo", "--run-sql", "--prefix", "ads_"])
_ex = [s for s, _ in rec.get("execute", [])]
check("Z6 --prefix 改写脚本里的表名与所有引用",
      "ods_t_order→ads_ods_t_order（1 处）" in out and "ods_t_user→ads_ods_t_user（3 处）" in out
      and all("ods_t_order" not in s.replace("ads_ods_t_order", "") for s in _ex), out[:200])
code, out, rec, _ = cli_fake([sql_a, "-d", "demo", "--run-sql", "--prefix", "ods_"])
check("Z6b 已带该前缀的表不叠加（重复跑是幂等的）",
      "未改写 2 张" in out and not any("ods_ods_" in s for s, _ in rec.get("execute", [])),
      out[:160])
code, out, rec, _ = cli_fake([sql_a, "-d", "demo", "--run-sql", "--swap-prefix", "dim"])
check("Z6c --swap-prefix 换掉脚本表名的首段（ods_t_order → dim_t_order）",
      any("CREATE TABLE `dim_t_order`" in s for s, _ in rec.get("execute", []))
      and "ods_t_order→dim_t_order" in out, out[:200])
code, out, rec, _ = cli_fake([sql_nohd, "-d", "demo", "--run-sql", "--swap-prefix", "dim"])
check("Z6d 表名里没有下划线的脚本不改，并逐条列出（与 GUI 同口径）",
      any("lonename" in s for s, _ in rec.get("execute", [])) and "跳过 1 张" in out, out[:200])

code, out, rec, _ = cli_fake([sql_a, "-d", "demo"])
check("Z7 不加 --run-sql 时 .sql 只登记、绝不连库",
      code == cli_mod.EXIT_OK and not rec and "脚本未执行" in out, (code, list(rec)))
code, out, rec, _ = cli_fake([sql_blank, "-d", "demo", "--run-sql"])
check("Z8 完全空白的脚本显式记为跳过，不开连接",
      "脚本为空" in out and not rec, (out[:120], list(rec)))
code, out, rec, _ = cli_fake([sql_empty, "-d", "demo", "--run-sql"])
check("Z8b 只有注释的脚本按 1 条语句发给 MySQL（MySQL 接受纯注释，不当空文件跳过）",
      len(rec.get("execute", [])) == 1 and "脚本为空" not in out,
      (len(rec.get("execute", [])), out[:120]))

code, out, rec, _ = cli_fake([sql_gbk, "-d", "demo", "--run-sql"])
check("Z9 GBK(记事本 ANSI) 存的脚本读得开、语句完整",
      any("gbk_one" in s for s, _ in rec.get("execute", [])), out[:150])

code, out, rec, _ = cli_fake([sql_bad, "-d", "demo", "--run-sql"],
                             fail_if=lambda sql: "NOT SQL" in sql.upper())
check("Z10 某条语句被拒：退出码 5、整文件回滚、第三条不再执行",
      code == cli_mod.EXIT_SCRIPT and rec.get("rollback") == 1 and rec.get("commit") is None
      and len(rec.get("execute", [])) == 2,
      (code, rec.get("rollback"), rec.get("commit"), len(rec.get("execute", []))))
check("Z10b 失败日志给出第几条、出错语句原文和回滚说明",
      "第 2/3 条语句被拒绝" in out and "本文件事务已回滚" in out
      and "出错语句" in out and "THIS IS NOT SQL" in out
      and not any("never_runs" in s for s, _ in rec.get("execute", [])),
      out[:200])

bad_dir = os.path.join(_tmpdir, "两个脚本")
os.makedirs(bad_dir, exist_ok=True)
shutil.copy(sql_bad, os.path.join(bad_dir, "坏语句.sql"))
shutil.copy(sql_a, os.path.join(bad_dir, "建表.sql"))
code, out, rec, _ = cli_fake([bad_dir, "-d", "demo", "--run-sql"],
                             fail_if=lambda sql: "NOT SQL" in sql.upper())
check("Z11 坏脚本不阻断同批其余脚本（成功 1 个 / 失败 1 个）",
      code == cli_mod.EXIT_SCRIPT and "成功 1 个文件，失败 1 个" in out
      and any("ods_t_order" in s for s, _ in rec.get("execute", [])), out[-200:])

code, out, _ = run_cli([sql_a, "-d", "demo", "--check", "--run-sql"])
check("Z12 --check 与 --run-sql 冲突时报用法错误（不静默忽略其中一个）",
      code == cli_mod.EXIT_USAGE and "参数冲突" in out, (code, out[:120]))
code, out, _ = run_cli([sql_a, "--run-sql"])
check("Z12b --run-sql 没给目标库直接判用法错（不会跑半截）",
      code == cli_mod.EXIT_USAGE, code)
code, out, _ = run_cli(["-d", "demo", "--run-sql"])
check("Z12c --run-sql 但输入里没有 .sql：明确警告而不是默默成功",
      code == cli_mod.EXIT_USAGE and "没有输入" in out, (code, out[:120]))

mix_dir = os.path.join(_tmpdir, "混合")
shutil.rmtree(mix_dir, ignore_errors=True)
os.makedirs(mix_dir)
shutil.copy(v_orders, os.path.join(mix_dir, "订单（ods_t_order）.xlsx"))
shutil.copy(sql_sel, os.path.join(mix_dir, "带查询.sql"))
code, out, rec, _ = cli_fake([mix_dir, "-d", "demo", "--run-sql"])
check("Z13 同一目录里 Excel 建表导入与脚本执行都跑（两条动作互不掺混）",
      code == cli_mod.EXIT_OK
      and any("CREATE TABLE `ods_t_order`" in s for s, _ in rec.get("execute", []))
      and any("CREATE TABLE `sel_one`" in s for s, _ in rec.get("execute", [])),
      (code, len(rec.get("execute", []))))
check("Z13b 混合跑的汇总里两边各有交代",
      "导入汇总" in out and "脚本汇总" in out, out[-220:])
code, out, rec, _ = cli_fake([mix_dir, "-d", "demo", "--run-sql"],
                             fail_if=lambda sql: "sel_one" in sql)
check("Z14 Excel 成功 + 脚本失败 → 取更严重的码 5",
      code == cli_mod.EXIT_SCRIPT and "导入汇总 | 成功 1 张表" in out, code)

code, out, rec, _ = cli_fake([os.path.join(_tmpdir, "没有这个.sql"), "-d", "demo", "--run-sql"])
check("Z15 不存在的 .sql 路径：明确报错并退回法码",
      code == cli_mod.EXIT_USAGE and "路径不存在" in out, (code, out[:120]))


# ---------------------------------------------------------------------------
section("G MySQL 真实读写（探针表跑完必须清零）")

G_LABELS = ("G1 真库连通 + 压缩包展开 + 导入退出码 0",
            "G1b 真库这一跑也没把密码写进日志",
            "G2 探针表真的建进库里（不是只没报错）",
            "G3 每张探针表都有数据",
            "G4 金额逐值回查：3 位小数没有被服务端圆成 2 位",
            "G5 日期列落库成 DATE（读回是 date 对象）",
            "G6 空单元格落成 NULL 而不是空串",
            "G7 表名带上前缀且分隔符没丢",
            "G8 探针表全部 DROP",
            "G9 库里其余表一张没少（没误伤别的表）",
            "G10 脚本执行到真库：退出码 0 并报出提交条数",
            "G11 脚本建的表真在库里",
            "G11b 脚本里的 INSERT 逐值落了库（含中文，不是只建了空表）")

PT = {k: os.environ.get("PT_" + k) for k in ("HOST", "PORT", "USER", "PASS", "DB")}
if not PT["USER"] or not PT["DB"]:
    for lbl in G_LABELS:
        skip(lbl, "未设 PT_USER/PT_DB（G 组要真连 MySQL，绝不冒充通过）")
else:
    PROBE = "e2m38probe"
    import pymysql

    def _db_tables(cur):
        cur.execute("SELECT TABLE_NAME FROM information_schema.TABLES "
                    "WHERE TABLE_SCHEMA=%s ORDER BY TABLE_NAME", (PT["DB"],))
        return [r[0] for r in cur.fetchall()]

    _c = pymysql.connect(host=PT["HOST"] or "127.0.0.1", port=int(PT["PORT"] or 3306),
                         user=PT["USER"], password=PT["PASS"], database=PT["DB"],
                         charset="utf8mb4", autocommit=True)
    with _c.cursor() as _cur:
        _baseline = _db_tables(_cur)
    _leftover = [x for x in _baseline if x.startswith(PROBE)]
    check("G0 开跑前探针表必须不存在（本组会 DROP 同名表）",
          not _leftover, _leftover or "干净")
    if not _leftover:
        g_name = "G探针（g_orders）.xlsx"
        g_src = make_xlsx(g_name, pd.DataFrame({
            "订单编号": ["P001", "P002", "P003"],
            "金额": [3.125, 1.5, 2.25],
            "下单日期": ["2026-04-01", "2026-04-02", "2026-04-03"],
            "备注": ["甲", "乙", None],
        }))
        g_dir = os.path.join(_tmpdir, "gprobe")
        os.makedirs(g_dir, exist_ok=True)
        g_zip = os.path.join(g_dir, "gprobe.zip")
        with zipfile.ZipFile(g_zip, "w") as zf:
            zf.write(g_src, g_name)
        os.remove(g_src)
        # 凭据走环境变量传给 CLI：命令行里不出现密码，日志里也不回显
        code, out, _ = run_cli([g_zip, "-d", PT["DB"], "--prefix", PROBE + "_"],
                               env={"E2M_HOST": PT["HOST"] or "127.0.0.1",
                                    "E2M_PORT": PT["PORT"] or "3306",
                                    "E2M_USER": PT["USER"], "E2M_PASS": PT["PASS"]})
        check("G1 真库连通 + 压缩包展开 + 导入退出码 0", code == cli_mod.EXIT_OK, (code, out[-400:]))
        check("G1b 真库这一跑也没把密码写进日志",
              bool(PT["PASS"]) and PT["PASS"] not in out, out[:120])
        with _c.cursor() as _cur:
            _now = _db_tables(_cur)
            _probes = [x for x in _now if x.startswith(PROBE)]
        check("G2 探针表真的建进库里（不是只没报错）", bool(_probes), _probes)
        _rows = {}
        for name in _probes:
            with _c.cursor() as _cur:
                _cur.execute("SELECT COUNT(*) FROM `%s`" % name)
                _rows[name] = _cur.fetchone()[0]
        check("G3 每张探针表都有数据", all(v > 0 for v in _rows.values()), _rows)
        _vals = []
        for name in _probes:
            with _c.cursor() as _cur:
                _cur.execute("SELECT `ding_dan_bian_hao`, `jin_e`, `xia_dan_ri_qi`, `bei_zhu` "
                             "FROM `%s` ORDER BY `ding_dan_bian_hao`" % name)
                _vals.append(_cur.fetchall())
        _flat = [r for rows in _vals for r in rows]
        check("G4 金额逐值回查：3 位小数没有被服务端圆成 2 位",
              any(str(r[1]) == "3.125" for r in _flat), _flat[:3])
        check("G5 日期列落库成 DATE（读回是 date 对象）",
              all(hasattr(r[2], "year") for r in _flat), _flat[:2])
        check("G6 空单元格落成 NULL 而不是空串",
              any(r[3] is None for r in _flat), _flat[:3])
        check("G7 表名带上前缀且分隔符没丢",
              all(x.startswith(PROBE + "_") for x in _probes), _probes)
        # 脚本线真跑：把一个 .sql 执行进真库，确认建表与 INSERT 都落下去
        g_sql = os.path.join(_tmpdir, "g_script.sql")
        with open(g_sql, "w", encoding="utf-8") as _fh:
            _fh.write("CREATE TABLE `%s_sql_orders` (\n"
                      "  `id` BIGINT NOT NULL,\n"
                      "  `name` VARCHAR(64) DEFAULT NULL COMMENT '名称',\n"
                      "  PRIMARY KEY (`id`)\n"
                      ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;\n"
                      "INSERT INTO `%s_sql_orders` VALUES (1, '甲'), (2, '乙');\n"
                      "SELECT `id` FROM `%s_sql_orders`;\n" % (PROBE, PROBE, PROBE))
        code, out, _ = run_cli([g_sql, "-d", PT["DB"], "--run-sql"],
                               env={"E2M_HOST": PT["HOST"] or "127.0.0.1",
                                    "E2M_PORT": PT["PORT"] or "3306",
                                    "E2M_USER": PT["USER"], "E2M_PASS": PT["PASS"]})
        check("G10 脚本执行到真库：退出码 0 并报出提交条数",
              code == cli_mod.EXIT_OK and "脚本提交" in out, (code, out[-200:]))
        with _c.cursor() as _cur:
            _sqltab = [x for x in _db_tables(_cur) if x == PROBE + "_sql_orders"]
            _rows = None
            if _sqltab:
                _cur.execute("SELECT `id`, `name` FROM `%s_sql_orders` ORDER BY `id`" % PROBE)
                _rows = _cur.fetchall()
        check("G11 脚本建的表真在库里", bool(_sqltab), _sqltab)
        check("G11b 脚本里的 INSERT 逐值落了库（含中文，不是只建了空表）",
              _rows == ((1, "甲"), (2, "乙")), _rows)

        # 收尾必须重新按前缀列一遍再 DROP：脚本那一步是中途新增的表，
        # 只清 _probes 那份早期快照会把 e2m38probe_sql_orders 留在库里
        _cleared = []
        with _c.cursor() as _cur:
            for name in [x for x in _db_tables(_cur) if x.startswith(PROBE)]:
                _cur.execute("DROP TABLE `%s`" % name)
                _cleared.append(name)
            _after = _db_tables(_cur)
        check("G8 探针表全部 DROP",
              not [x for x in _after if x.startswith(PROBE)], _cleared)
        check("G9 库里其余表一张没少（没误伤别的表）",
              len(_after) == len(_baseline), "基线 %d → 收尾 %d" % (len(_baseline), len(_after)))
    else:
        for lbl in G_LABELS:
            skip(lbl, "上一轮遗留了探针表 %s，本轮不动库，请先人工确认再删" % _leftover)
    _c.close()


# ---------------------------------------------------------------------------
print("\n" + "=" * 72)
print("用例集结束（%s · Python %s）：通过 %d  失败 %d  跳过 %d"
      % (STACK, sys.version.split()[0], RESULTS["pass"], RESULTS["fail"], RESULTS["skip"]))
if FAILURES:
    print("\n失败用例：")
    for f in FAILURES:
        print("  - " + f)
if SKIPPED:
    print("\n未覆盖（显式记账，不算通过）：")
    for label, why in SKIPPED:
        print("  - %s：%s" % (label, why))
print("=" * 72)
SUMMARY_PRINTED = True
shutil.rmtree(_tmpdir, ignore_errors=True)
sys.exit(1 if RESULTS["fail"] else 0)
