# -*- coding: utf-8 -*-
"""把 Python(PySide6) 那条 GUI 线的界面与行为枚举成扁平 JSON，供与 Swing 线逐键对撞。

为什么不用"读代码列条目"那套：上一轮就是这么自证的，结果用户连说五次"差距还很多"。
清单是自己写的，边界也就停在自己想到的地方。这里改成**从活对象上机器枚举**：
每个按钮/输入框/表格/弹窗都按 key 落一份事实，两条线的 key 集与值由比较器逐条对，
少一个按钮、文案差一个字、启用条件不同，都会以"键缺失/值不等"的形式露出来。

用法（仓库根目录）：
    .venv/Scripts/python.exe java/tools/gui_dump_py.py <输出json> [夹具目录]
"""
import json
import os
import re
import shutil
import sys
import tempfile

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    tempfile.gettempdir(), "e2m-gui-dump", "py.json")
FIX = sys.argv[2] if len(sys.argv) > 2 else os.path.join(REPO, "java", "verify", "fixtures")

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, REPO)

from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtWidgets import (QApplication, QCheckBox, QComboBox, QGridLayout, QGroupBox,  # noqa: E402
                               QLabel, QLineEdit, QListWidget, QPlainTextEdit, QProgressBar,
                               QPushButton, QScrollArea, QSpinBox, QTableWidget)
import main as m  # noqa: E402

app = QApplication(sys.argv)
m.setup_app(app)   # 图标等应用级设置走生产同一段代码
m.MODAL_PROGRESS = False
POPUPS = []


class _StubBox(object):
    def __init__(self, parent=None):
        self.parent = parent
        self._kind = "information"
        self._title = ""
        self._text = ""

    def setWindowTitle(self, title):
        self._title = title

    def setText(self, text):
        self._text = text

    def setIcon(self, icon):
        name = getattr(icon, "name", str(icon))
        self._kind = {"Information": "information", "Warning": "warning",
                      "Critical": "critical"}.get(name, "information")

    def setStandardButtons(self, buttons):
        pass

    def findChildren(self, cls):
        return []

    def exec(self):
        POPUPS.append((self._kind, self._title, self._text))
        return 0

    @staticmethod
    def information(parent, title, text, *a, **k):
        POPUPS.append(("information", title, text))
        return 0

    @staticmethod
    def warning(parent, title, text, *a, **k):
        POPUPS.append(("warning", title, text))
        return 0

    @staticmethod
    def critical(parent, title, text, *a, **k):
        POPUPS.append(("critical", title, text))
        return 0

    @staticmethod
    def question(parent, title, text, *a, **k):
        POPUPS.append(("question", title, text))
        return 0


m.QMessageBox = _StubBox
m._POPUP_MESSAGE_BOX = _StubBox
os.chdir(REPO)

D = {}
TMP = tempfile.mkdtemp(prefix="e2m-gui-dump-")
_PATH_RE = re.compile(r"[A-Za-z]:[\\/][^|;\"'\]\[）)）]*")
_DAY_RE = re.compile(r"(meta_data_file_|simple_meta_data_file_)\d{8}")


def norm(text):
    s = "" if text is None else str(text)
    s = s.replace("\r", "")
    s = _PATH_RE.sub(lambda mo: "<P>" + os.path.basename(mo.group(0).rstrip("/\\")), s)
    s = _DAY_RE.sub(lambda mo: mo.group(1) + "<DAY>", s)
    return s.strip()


def put(key, value):
    D[key] = norm(value)


def tip_of(w):
    try:
        return w.toolTip()
    except Exception:
        return ""


def dump_widget(role, w):
    if w is None:
        put("w.%s.missing" % role, "1")
        return
    put("w.%s.kind" % role, type(w).__name__)
    if not isinstance(w, (QSpinBox, QComboBox)):
        put("w.%s.text" % role, getattr(w, "text")() if hasattr(w, "text") else "")
    put("w.%s.tip" % role, tip_of(w))
    put("w.%s.on" % role, int(bool(w.isEnabled())))
    put("w.%s.vis" % role, int(bool(w.isVisible())))
    if isinstance(w, QLineEdit):
        put("w.%s.ph" % role, w.placeholderText())
        put("w.%s.echo" % role, w.echoMode().name)
    if isinstance(w, QSpinBox):
        put("w.%s.value" % role, w.value())
        put("w.%s.min" % role, w.minimum())
        put("w.%s.max" % role, w.maximum())
    if isinstance(w, QCheckBox):
        put("w.%s.checked" % role, int(bool(w.isChecked())))
    if isinstance(w, QProgressBar):
        put("w.%s.value" % role, w.value())
        put("w.%s.min" % role, w.minimum())
        put("w.%s.max" % role, w.maximum())
        put("w.%s.text" % role, w.format().replace("%p", "0"))
    if isinstance(w, QComboBox):
        put("w.%s.items" % role, ",".join(w.itemText(i) for i in range(w.count())))


def page_of(win, w):
    """按钮属于哪一页：两页都有「批量应用前缀」「清除日志」，不分页就会互相覆盖。

    QTabWidget 的页面控件挂在内部 QStackedWidget 下，不是 main_tabs 的直接子控件，
    所以只能自下而上找"哪一页控件是我的祖先"。
    """
    tabs = win.main_tabs
    for i in range(tabs.count()):
        page = tabs.widget(i)
        cur = w
        while cur is not None:
            if cur is page:
                return norm(tabs.tabText(i))
            cur = cur.parent()
    return "顶部"


def dump_buttons(win):
    seen = []
    for b in win.findChildren(QPushButton):
        t = norm(b.text())
        if not t:
            continue
        page = page_of(win, b)
        key = "btn.%s.%s" % (page, t)
        seen.append("%s|%s" % (page, t))
        put(key + ".on", int(bool(b.isEnabled())))
        put(key + ".tip", tip_of(b))
    put("btn.list", "|".join(sorted(seen)))
    bar = find_layout(win.main_tabs.widget(0), win.btn_import)
    order = []
    if bar is not None:
        for i in range(bar.count()):
            wd = bar.itemAt(i).widget()
            if wd is not None and isinstance(wd, QPushButton) and norm(wd.text()):
                order.append(norm(wd.text()))
    put("bar.excel.order", "|".join(order))
    dump_sql_rows(win)


def title_of_table(table):
    """这张表格左上角那句标识：从真控件往上找到那层组框，读它自己的标题。

    对撞器不能两边各抄一份常量——抄来的那份改了，另一份不会红。
    """
    cur = table
    while cur is not None:
        if isinstance(cur, QGroupBox):
            return norm(cur.title())
        cur = cur.parentWidget()
    return ""


def dump_sql_rows(win):
    """第二页每一行从左到右是哪几颗按钮——位置也是界面事实，与 Java 侧同判据。

    按钮文案早就在逐颗比，但"清空待上传文件排在开关那一行、还是排在保存/执行那一行"
    只有看真正的行结构才比得出来。按像素量不行：这一页在没被切到前面时子控件还没排过版，
    量出来全是 0，所以直接读布局——表格区那层布局里的每一条 QHBoxLayout 就是一行。
    """
    upper = None
    for i in range(win.sql_split.count()):
        w = win.sql_split.widget(i)
        # 日志段现在套了一层面板（清除日志在框上方），按里面那个文本框认
        if w is not win.sql_drop and w.findChild(QPlainTextEdit) is not win.sql_log:
            upper = w
    if upper is None or upper.layout() is None:
        return
    out = []
    band = 0
    for j in range(upper.layout().count()):
        item = upper.layout().itemAt(j)
        lay = item.layout()
        if lay is None:          # 第一个 item 是表格本身，不是行
            continue
        names = []
        for k in range(lay.count()):
            wd = lay.itemAt(k).widget()
            if isinstance(wd, QPushButton) and norm(wd.text()):
                names.append(norm(wd.text()))
        if names:
            out.append("%d:%s" % (band, ",".join(names)))
            band += 1
    # 底部动作条（保存/执行/生成）与覆盖开关已挪到分隔条外，和第一页同一形状；
    # 它不是表格区里的一行，但“哪几颗按钮在同一排”仍是界面事实，按下一带接着比
    bar = find_layout(win.main_tabs.widget(1), win.btn_sql_save)
    if bar is not None:
        names = []
        for k in range(bar.count()):
            wd = bar.itemAt(k).widget()
            if isinstance(wd, QPushButton) and norm(wd.text()):
                names.append(norm(wd.text()))
        if names:
            out.append("%d:%s" % (band, ",".join(names)))
    put("bar.sql.rows", "|".join(out))


def dump_table(role, tw):
    if not isinstance(tw, QTableWidget):
        return
    cols = []
    for c in range(tw.columnCount()):
        it = tw.horizontalHeaderItem(c)
        cols.append(norm(it.text()) if it else "")
    put("table.%s.cols" % role, "|".join(cols))
    put("table.%s.width" % role, tw.columnCount())
    put("table.%s.zebra" % role, int(bool(tw.alternatingRowColors())))
    rows = []
    for r in range(tw.rowCount()):
        cells = []
        for c in range(tw.columnCount()):
            it = tw.item(r, c)
            cells.append(norm(it.text()) if it else "")
        rows.append(";".join(cells))
    put("table.%s.rows" % role, "\n".join(rows))


def find_layout(root, target):
    """在 root 这棵子树里找真正装着 target 的那层布局。

    Qt 的嵌套 layout 不会把控件改父级，parentWidget().layout() 拿到的往往是整页的
    QVBoxLayout；而页与页之间是 widget 项而不是 layout 项，所以两样都要往下递归。
    """
    stack = [root.layout()] if root.layout() is not None else []
    while stack:
        lay = stack.pop()
        for i in range(lay.count()):
            item = lay.itemAt(i)
            if item.widget() is target:
                return lay
            if item.layout() is not None:
                stack.append(item.layout())
            w = item.widget()
            if w is not None and w.layout() is not None:
                stack.append(w.layout())
            # 横向滚动区里的内容挂在自己的滚动视口里，不是它的直接 layout；
            # 底部动作条就套了这一层，不穿过去会读成"这排按钮不存在"
            if isinstance(w, QScrollArea) and w.widget() is not None \
                    and w.widget().layout() is not None:
                stack.append(w.widget().layout())
    return None


def popup_body(text):
    """弹窗正文去掉换行，只留内容（与 Java 侧 popupBody 同一件事）。

    这里的换行就是写文案的人自己分的段（正文不再被折行改写），
    比"说的话"时不连分段一起比。
    """
    return (text or "").replace("\n", "")


def popup_text():
    """最近一次弹窗的"图标|标题|正文"，没弹就是空串。

    两边显示的是同一份原文：横向折行交给各自的控件按内容区右边界做，
    所以这里只比"说的话"。排版本身由 popup.display.raw 那一把键钉住
    （显示的必须就是原文，谁都不许预先往里插换行）。
    """
    if not POPUPS:
        return ""
    _kind, title, text = POPUPS[-1]
    return "%s|%s|%s" % (_kind, title, text.replace("\n", ""))


def _db_tone(w):
    """目标库信息条现在是哪一档颜色：绿色=已选定、红色=没选库，别的写法算 none。"""
    s = w.styleSheet()
    if "#1a7f37" in s:
        return "ok"
    if "#c62828" in s:
        return "bad"
    return "none"


def btn_on(win, text):
    """按文案取「Excel 导入」页上那颗按钮的当前可点状态（两页重名的按钮按页取）。"""
    for b in win.findChildren(QPushButton):
        if norm(b.text()) == text and page_of(win, b) == u"Excel 导入":
            return int(bool(b.isEnabled()))
    return "缺"


def names(win):
    return "|".join(t["table_name"] for t in win.tables)


def main():
    # 对撞口径固定在"没有 conn.ini"：否则程序同级那份配置会触发启动自动连接，
    # 探针结果随运行目录、环境变量变（Java 侧 GuiDump 同样钉住）
    no_cfg = os.path.join(tempfile.gettempdir(), "e2m-gui-dump-python-noconf")
    os.makedirs(no_cfg, exist_ok=True)
    for k in [k for k in os.environ if k.startswith("E2M_")]:
        del os.environ[k]
    keep_pdir = m.program_dir
    m.program_dir = lambda: no_cfg
    win = m.MainWindow()
    win.show()
    app.processEvents()
    m.program_dir = keep_pdir

    put("tabs", "|".join(win.main_tabs.tabText(i)
                        for i in range(win.main_tabs.count())))
    dump_widget("host", win.ed_host)
    dump_widget("port", win.ed_port)
    dump_widget("user", win.ed_user)
    dump_widget("pass", win.ed_pass)
    dump_widget("db", win.cb_db)
    dump_widget("prefix", win.ed_prefix)
    dump_widget("sql_prefix", win.ed_sql_prefix)
    dump_widget("drop_existing", win.chk_drop)
    dump_widget("sql_drop_existing", win.chk_sql_drop)
    dump_widget("initials", win.chk_initials)
    dump_widget("upper", win.chk_upper)
    dump_widget("sql_upper", win.chk_sql_upper)
    dump_widget("progress", win.progress)
    dump_buttons(win)
    dump_table("file", win.file_table)
    dump_table("sql", win.sql_table)
    # 每张主表格左上角那句标识（读真控件上的组框标题，不是各抄一份常量）
    put("table.file.title", title_of_table(win.file_table))
    put("table.sql.title", title_of_table(win.sql_table))
    # 字号也是界面事实：比的是设计稿那一档（pt），不是各台机器上的物理像素
    put("font.table.pt", str(win.file_table.font().pointSize()))
    put("font.log.pt", str(win.log.font().pointSize()))
    # 连接区：跑业务时收起成一条状态栏，只报当前是哪个库（用户口径）
    put("conn.state", win.lbl_conn_state.text())
    put("conn.expanded", "展开" if win.conn_form.isVisible() else "收起")
    put("conn.edit", "显示" if win.btn_conn_edit.isVisible() else "隐藏")
    put("hint.excel", win.lbl_imp_db.text())
    put("hint.sql", win.lbl_sql_db.text())
    put("intro.excel", norm(win.imp_intro.text()))
    put("intro.sql", norm(win.sql_intro.text()))
    put("drop.hint", win.drop.label.text())
    put("drop.sql.hint", win.sql_drop.label.text())
    put("drop.sql.accepts", "true" if win.sql_drop.acceptDrops() else "false")
    # 拖拽区的浅色底与悬停高亮：读真控件存的色值，两栈同值才算同一块区域
    put("drop.fill", win.drop.fill_color)
    put("drop.fill.hover", win.drop.hover_color)
    put("drop.sql.fill", win.sql_drop.fill_color)
    put("drop.sql.fill.hover", win.sql_drop.hover_color)
    put("sqlrun.drop.fill", win.sql_run_drop.fill_color)
    put("sqlrun.drop.fill.hover", win.sql_run_drop.hover_color)
    # 空清单时表格里的引导语与上方虚线区呼应
    put("table.file.empty", win.file_table.empty_hint)
    put("table.sql.empty", win.sql_table.empty_hint)
    put("table.sqlrun.empty", win.sql_run_table.empty_hint)
    # 底部动作条底板：覆盖开关与主按钮是否在同一块浅色底里（1|1 才算归组）
    put("action.excel.bg", win.excel_action_bar.fill_color)
    put("action.excel.group", "%d|%d" % (
        int(win.excel_action_bar.isAncestorOf(win.chk_drop)),
        int(win.excel_action_bar.isAncestorOf(win.btn_import))))
    put("action.sql.bg", win.sql_action_bar.fill_color)
    put("action.sql.group", "%d|%d" % (
        int(win.sql_action_bar.isAncestorOf(win.chk_sql_drop)),
        int(win.sql_action_bar.isAncestorOf(win.btn_sql_run))))
    # 清单空时两页的动作按钮都要在既有禁用样式上灰着：首屏对撞的就是这个状态
    put("action.excel.enabled", "%d|%d" % (
        int(win.btn_gen_std.isEnabled()), int(win.btn_gen_simple.isEnabled())))
    put("action.sql.enabled", "%d|%d|%d|%d" % (
        int(win.btn_sql_save.isEnabled()), int(win.btn_sql_run.isEnabled()),
        int(win.btn_sql_gen_std.isEnabled()), int(win.btn_sql_gen_simple.isEnabled())))
    put("action.run.bg", win.run_action_bar.fill_color)
    put("action.run.group", "%d|%d" % (
        int(win.run_action_bar.isAncestorOf(win.chk_run_drop)),
        int(win.run_action_bar.isAncestorOf(win.btn_run_scripts))))

    def _sql_segment(w):
        if w is win.sql_drop:
            return u"拖拽区"
        if w.findChild(QTableWidget) is win.sql_table:
            return u"表格区"
        # 日志段现在套了一层面板（清除日志在框上方），按里面那个文本框认
        if w.findChild(QPlainTextEdit) is win.sql_log:
            return u"日志区"
        return type(w).__name__

    put("split.sql.order", "|".join(_sql_segment(win.sql_split.widget(i))
                                    for i in range(win.sql_split.count())))

    # 日志默认高度 = 它自己的最小高度（用户口径：默认压到最低，仍可拖）。
    # 量的是日志「段」（含上方清除日志那排），与用例/自检同一口径；不是里面那个文本框。
    def _log_at_min(panel):
        minimum = max(panel.minimumHeight(), panel.minimumSizeHint().height())
        return "1" if 0 <= panel.height() - minimum <= 4 else "0"

    put("split.excel.logmin", _log_at_min(win.excel_split.widget(2)))
    win.main_tabs.setCurrentIndex(1)
    app.processEvents()
    put("split.sql.logmin", _log_at_min(win.sql_split.widget(2)))

    # 第三页默认把「编辑区与结果」压到最小，高度让给脚本清单：量真控件，
    # 行数键跨栈同值（两边的字体像素不同，比行数不比像素）
    def _run_segment(w):
        if w is win.sql_run_drop:
            return u"拖拽区"
        if w.findChild(QTableWidget) is win.sql_run_table:
            return u"脚本清单"
        if w is win.sql_run_split:
            return u"编辑区与结果"
        return type(w).__name__

    win.main_tabs.setCurrentIndex(2)
    app.processEvents()
    put("split.sqlrun.order", "|".join(
        _run_segment(win.sql_run_split_v.widget(i))
        for i in range(win.sql_run_split_v.count())))
    put("split.sqlrun.tailmin", _log_at_min(win.sql_run_split))
    put("sqlrun.editor.minlines", str(int(round(
        win.sql_run_input.minimumHeight()
        / float(win.sql_run_input.fontMetrics().height())))))
    put("sqlrun.table.fits", "1" if win.sql_run_table.viewport().height()
        >= 44 + 10 + win.sql_run_table.fontMetrics().height() else "0")
    win.main_tabs.setCurrentIndex(0)
    app.processEvents()

    # ---- 第三个页签「SQL 执行器」：常驻可见量 + 示例对话框的一来一回 ----
    put("hint.sqlrun", norm(win.sql_run_hint.text()))
    put("hint.sqlrun.db", win.lbl_run_db.text())
    # 第三页的脚本队列：拖拽区、清单表、覆盖开关与底部那一排按钮
    put("sqlrun.drop.hint", win.sql_run_drop.label.text())
    put("sqlrun.drop.accepts", "true" if win.sql_run_drop.acceptDrops() else "false")
    dump_widget("run_drop_existing", win.chk_run_drop)
    dump_table("sqlrun", win.sql_run_table)
    put("table.sqlrun.title", title_of_table(win.sql_run_table))
    put("sqlrun.noedit", "1" if win.sql_run_table.editTriggers()
        == m.QAbstractItemView.NoEditTriggers else "0")
    _run_bar = find_layout(win.main_tabs.widget(2), win.btn_run_scripts)
    _run_names = []
    if _run_bar is not None:
        for i in range(_run_bar.count()):
            wd = _run_bar.itemAt(i).widget()
            if isinstance(wd, QPushButton) and norm(wd.text()):
                _run_names.append(norm(wd.text()))
    put("bar.sqlrun.order", "|".join(_run_names))
    # 页面上的静态标签按字典序连起来比：两边的控件遍历顺序本来就不同（Qt 按创建序、
    # Swing 按加进容器的序），排序后比的是"这一页上有哪几句话"。拖拽区自己那句提示
    # 是画出来的（Java 侧同样直接画在画布上），不算这一页的静态标签。
    put("sqlrun.labels", "|".join(sorted(
        norm(l.text()) for l in win.main_tabs.widget(2).findChildren(QLabel)
        if norm(l.text()) and not win.sql_run_drop.isAncestorOf(l))))
    put("sqlrun.default", win.sql_run_input.toPlainText())
    put("sqlrun.wrap", "wrap" if win.sql_run_input.lineWrapMode()
        == QPlainTextEdit.WidgetWidth else "nowrap")
    POPUPS.clear()
    win.sql_run_input.clear()
    win.on_sql_run()
    put("guard.sqlrun.empty", popup_text())
    win.sql_run_input.setPlainText(m.SQL_RUN_DEFAULT)
    win.on_sql_samples()
    app.processEvents()
    _sd = win.sql_run_samples
    put("samples.title", _sd.windowTitle())
    put("samples.count", _sd.list.count())
    put("samples.item.0", norm(_sd.list.item(0).text()))
    put("samples.preview.0", _sd.preview.toPlainText())
    # 全文预览的字号也是界面事实：这是与左侧清单并排的阅读面，不吃编辑区那一档（编辑区已随日志档）
    put("samples.preview.font.pt", str(_sd.preview.font().pointSize()))
    _sd.list.setCurrentRow(7)
    put("samples.preview.7", _sd.preview.toPlainText())
    # 不比对「复制到剪贴板」的那句回执：剪贴板可能被别的进程占用，
    # 两条线各跑一次的成败不同，那是现场状态不是界面差异（两条线各自的自检里咬回执本身）
    _sd.show()
    win.sql_run_input.setPlainText(m.SQL_RUN_DEFAULT)
    _sd.list.setCurrentRow(3)
    _sd.btn_insert.click()
    put("samples.insert.default", win.sql_run_input.toPlainText())
    _sd.show()
    win.sql_run_input.setPlainText("SELECT 1")
    _sd.list.setCurrentRow(1)
    _sd.btn_insert.click()
    # 追加那条只比"补没补分号换行"：两条线的行分隔符不同（\r\n 与 \n），
    # 正文本身已由 samples.sql.1 那一键钉住。
    put("samples.insert.append",
        "yes" if win.sql_run_input.toPlainText().startswith("SELECT 1;\n") else "no")
    win.sql_run_input.setPlainText(m.SQL_RUN_DEFAULT)

    # ---- 空态：哪些按钮该是灰的 ----
    for t in ("开始导入", "生成标准excel文件", "生成简易excel文件", "预览 DDL",
              "生成表结构预览", "停止"):
        put("empty.on.%s" % t, btn_on(win, t))

    # ---- 空列表 / 没选库 / 拖错文件 / 空点停止：把说的那句话取出来对 ----
    POPUPS.clear()
    win.on_preview_ddl()
    put("guard.empty.ddl", popup_text())
    POPUPS.clear()
    win.on_preview_structure()
    put("guard.empty.structure", popup_text())
    POPUPS.clear()
    win.on_gen_standard()
    put("guard.empty.gen_std", popup_text())
    POPUPS.clear()
    win.on_gen_simple()
    put("guard.empty.gen_simple", popup_text())
    mark = win.log.toPlainText()
    win.on_files_dropped([os.path.join(TMP, "readme.txt")])
    put("guard.drop.txt", win.log.toPlainText()[len(mark):].strip())
    mark = win.log.toPlainText()
    win.on_stop()
    put("guard.stop", win.log.toPlainText()[len(mark):].strip())

    # ---- 连不上库时的那句话：只比固定文案那一行 ----
    # 驱动细节（pymysql 的 2003 与 JDBC 的 Communications link failure）必然不同，
    # 那是运行时差异不是界面差异；但"无法连接到 主机:端口（用户 X），库列表未加载"
    # 这一句和弹窗标题是两条线共用的契约。
    win.ed_host.setText("127.0.0.1")
    win.ed_port.setValue(1)
    POPUPS.clear()
    mark = win.log.toPlainText()
    win.on_load_databases()
    put("guard.conn.popup", "%s|%s" % (POPUPS[-1][0], POPUPS[-1][1]) if POPUPS else "")
    put("guard.conn.log", win.log.toPlainText()[len(mark):].strip().split("\n")[0])
    put("guard.conn.db_enabled", int(bool(win.cb_db.isEnabled())))

    # ---- 加载夹具 ----
    files = sorted(os.path.join(FIX, f) for f in os.listdir(FIX)
                   if f.lower().endswith((".xlsx", ".xls")))
    win._add_files(files)
    app.processEvents()
    put("loaded.names", names(win))
    put("file.serial", "%s=%s" % (
        win.file_table.horizontalHeaderItem(0).text(),
        ",".join(win.file_table.item(i, 0).text()
                 for i in range(win.file_table.rowCount()))))
    dump_table("file", win.file_table)
    # 两把与 Del 键有关的跨线事实：按整行选中（不这样，"删这一行"在两边语义就不同）；
    # 以及空选时按下 Del 的那句留痕文案。都只动选中态、不改清单，取完复原。
    from PySide6.QtCore import QEvent as _Ev
    from PySide6.QtGui import QKeyEvent as _Kv
    win.file_table.setCurrentCell(1, 3)
    app.processEvents()
    _sel_rows = [i.row() for i in win.file_table.selectionModel().selectedRows()]
    put("file.rowselect", "%d/%d" % ((_sel_rows or [-1])[0], len(_sel_rows)))
    win.file_table.clearSelection()
    win.file_table.keyPressEvent(_Kv(_Ev.KeyPress, Qt.Key_Delete, Qt.NoModifier))
    app.processEvents()
    _log_tail = win.log.toPlainText()
    _at = _log_tail.rfind(u"[删除]")
    put("file.del.empty", _log_tail[_at:].splitlines()[0] if _at >= 0 else u"(没留痕)")
    win.file_table.clearSelection()
    for t in ("开始导入", "生成标准excel文件", "生成简易excel文件"):
        put("loaded.on.%s" % t, btn_on(win, t))
    put("loaded.log", win.log.toPlainText())
    # 日志框最小高度按"行数"约定（两栈同一档字号的行高不同，比像素数会各说各话）：
    # 缩成两行的日志框等于没有，这条比的就是两边都留够四行
    put("log.excel.minh", str(m.LOG_MIN_LINES))
    put("log.sql.minh", str(m.LOG_MIN_LINES))
    put("log.run.minh", str(m.LOG_MIN_LINES))
    put("loaded.hint.excel", win.lbl_imp_db.text())
    put("add.popup", popup_text())
    # 两侧都量"程序自己画的那块地方"：Qt 顶层窗口的 width/height 就是客户区。
    # 跨线比的是配置值 want（两边必须写同一个数）；实测值受现场屏幕与夹取影响，只记不比。
    put("window.want", str(win.property("e2mWant") or ""))
    put("env.window.size", "%dx%d" % (win.width(), win.height()))
    ic = app.windowIcon()   # 图标挂在 QApplication 上（main.py 末尾 setWindowIcon），不是挂在窗口上
    sizes = sorted(ic.availableSizes(), key=lambda s: max(s.width(), s.height()))
    put("window.icon", "none" if not sizes
        else "max=%d" % max(sizes[-1].width(), sizes[-1].height()))
    tip0 = win.file_table.item(0, 1)
    put("file.tooltip.0", tip0.toolTip() if tip0 else "")
    # 文件表只有「表名」列可编辑；单击改名、非法名恢复原名，两条线逐句话比
    put("file.rename.editable", "|".join(
        str(int(bool(win.file_table.item(0, c).flags() & Qt.ItemIsEditable)))
        for c in (2, 0, 3)))
    _rename_old = win.tables[0]["table_name"]
    win.file_table.item(0, 2).setText(_rename_old + "_probe")
    app.processEvents()
    _log_lines = [ln for ln in win.log.toPlainText().splitlines() if ln.strip()]
    put("file.rename.ok", "%s|%s|%s" % (
        win.tables[0]["table_name"], win.tables[0]["base_table_name"],
        _log_lines[-1] if _log_lines else ""))
    POPUPS.clear()
    win.file_table.item(0, 2).setText("")
    app.processEvents()
    put("file.rename.empty", "%s|%s" % (win.tables[0]["table_name"], popup_text()))
    POPUPS.clear()
    win.file_table.item(0, 2).setText("bad name")
    app.processEvents()
    put("file.rename.bad", "%s|%s" % (win.tables[0]["table_name"], popup_text()))
    POPUPS.clear()
    win.file_table.item(0, 2).setText(win.tables[1]["table_name"])
    app.processEvents()
    put("file.rename.dupe", "%s|%s" % (win.tables[0]["table_name"], popup_text()))
    win.file_table.item(0, 2).setText(_rename_old)
    app.processEvents()
    put("file.rename.restored", "%s|%s" % (
        win.tables[0]["table_name"], win.tables[0]["base_table_name"]))
    # 弹窗显示的正文必须就是原文：折行是排版的活，谁都不许预先往文本里插换行
    _worst = ("任务异常：读取 C:/数据/很长的一个文件名防空警报试鸣信息（dim_air_raid_siren）.xlsx 时失败，"
              "内部单元格 D17 的类型与推断结果冲突，需要人工确认这一列到底是文本还是数字，"
              "详见日志里的完整堆栈信息。")
    _raw_box = m.build_popup_box(None, "任务异常", _worst, "critical",
                                 box_class=m.PopupMessageBox)
    put("popup.display.raw", "原文" if _raw_box.text() == _worst else "被改写过")
    _raw_box.deleteLater()
    # 弹窗按正文给尺寸：三把键钉住这条规则的两端（都落在夹取边界上，
    # 不受两边字体度量差几个像素的影响），中间值由两边的用例各自量。
    # 长句那一条只有宽度落在边界上，高度按折完的行数走，两边字体不同，只比宽度
    for _case, _txt, _full in (("short", "清单里没有脚本文件，无需清空", True),
                               ("tall", "\n".join("第 %d 行报错正文。" % i
                                                  for i in range(1, 41)), True),
                               ("longwidth", _worst, False)):
        _sbox = m.build_popup_box(None, "提示", _txt, "warning",
                                  box_class=m.PopupMessageBox)
        _w = str(_sbox.property("e2mWant") or "")
        put("popup.want." + _case, _w if _full else _w.split("x")[0])
        _sbox.deleteLater()
    # 双击不是"接了信号"就算数，要真的建出那一张表的结构窗口。
    # 真窗口是模态的，离屏 exec() 会挂住进程，所以只把 exec 换成"记一笔、返回取消"，
    # 构造过程走真代码。
    _real_sd = m.StructureDialog
    _opened = []

    class _ProbeSd(_real_sd):
        def exec(self):
            _opened.append(self)
            return 0

    m.StructureDialog = _ProbeSd
    try:
        win.file_table.setCurrentCell(0, 0)
        win.file_table.cellDoubleClicked.emit(0, 0)
    finally:
        m.StructureDialog = _real_sd
    if _opened:
        _sd = _opened[0]
        _titles = "|".join(_sd.tabs.tabText(i) for i in range(_sd.tabs.count()))
        _sd.deleteLater()
        put("dblclick.structure", "1|%d|%s" % (_sd.tabs.count(), _titles))
    else:
        put("dblclick.structure", "0|0|")
    POPUPS.clear()
    win.on_start_import()
    put("guard.nodb.import", popup_text())

    # ---- 结构弹窗改一格 → DDL 要跟着变（编辑真的落进表定义）----
    _t0 = win.tables[0]["columns"][0]
    _before_type, _before_ddl = _t0["type"], None
    _t0["type"] = "VARCHAR(123)"
    _dlg = m.DDLDialog(win.tables, False, win)
    put("structure.edit.ddl", "1" if "VARCHAR(123)" in _dlg.tabs.widget(0).toPlainText() else "0")
    _dlg.deleteLater()
    _t0["type"] = _before_type

    # ---- 压缩包：展开、汇总、来源 tooltip、临时目录清理 ----
    zdir = os.path.join(TMP, "zip")
    os.makedirs(zdir)
    import zipfile as _zf
    zpath = os.path.join(zdir, "pack1.zip")
    with _zf.ZipFile(zpath, "w") as zf:
        zf.writestr("readme.txt", "不是 Excel")
        with open(files[0], "rb") as fh:
            zf.writestr("inner.xlsx", fh.read())
    _n_before = len(win.tables)
    POPUPS.clear()
    mark = win.log.toPlainText()
    win._add_files([zpath])
    app.processEvents()
    put("zip.added", len(win.tables) - _n_before)
    put("zip.log", win.log.toPlainText()[len(mark):].strip())
    put("zip.popup", popup_text())
    _last = win.file_table.rowCount() - 1
    _it = win.file_table.item(_last, 1)
    put("zip.tooltip", _it.toolTip() if _it else "")

    # ---- 前缀两步 ----
    win.ed_prefix.setText("ods_")
    win.on_apply_prefix()
    app.processEvents()
    put("prefix.ods.names", names(win))
    win.ed_prefix.setText("change")
    win.on_replace_prefix()
    app.processEvents()
    put("prefix.change.names", names(win))
    win.ed_prefix.setText("")
    win.on_apply_prefix()
    app.processEvents()
    put("prefix.cleared.names", names(win))

    # ---- DDL：不勾选与勾选覆盖已存在表 ----
    dlg = m.DDLDialog(win.tables, False, win)
    put("dlg.ddl.want", str(dlg.property("e2mWant") or ""))
    put("env.dlg.ddl.size", "%dx%d" % (dlg.width(), dlg.height()))
    put("ddl.tabs", "|".join(dlg.tabs.tabText(i) for i in range(dlg.tabs.count())))
    put("ddl.first", dlg.tabs.widget(0).toPlainText())
    dlg.deleteLater()
    win.chk_drop.setChecked(True)
    dlg2 = m.DDLDialog(win.tables, True, win)
    put("ddl.drop.first", dlg2.tabs.widget(0).toPlainText())
    dlg2.deleteLater()
    win.chk_drop.setChecked(False)

    # ---- 结构弹窗 ----
    sd = m.StructureDialog(win.tables[:2], win)
    put("dlg.structure.title", sd.windowTitle())
    put("dlg.structure.want", str(sd.property("e2mWant") or ""))
    put("env.dlg.structure.size", "%dx%d" % (sd.width(), sd.height()))
    put("dlg.structure.tabs", "|".join(sd.tabs.tabText(i)
                                       for i in range(sd.tabs.count())))
    first = sd.tabs.widget(0).findChild(QTableWidget)
    put("dlg.structure.cols", "|".join(norm(first.horizontalHeaderItem(c).text())
                                       for c in range(first.columnCount())))
    # 第 4 列是实测字节数（只读）、页上方那颗输入框是表注释（可改）：
    # 这两件事两边各写一份，能不能改必须逐列对，不看就只能靠肉眼
    put("dlg.structure.comment", norm(sd.comment_edits[0].text()))
    put("dlg.structure.editable", "|".join(
        "1" if (first.item(0, c) is not None and
                 bool(first.item(0, c).flags() & Qt.ItemIsEditable)) else "0"
        for c in range(1, 5)))

    def srow(r):
        cells = []
        for c in range(first.columnCount()):
            if c == 0:
                cw = first.cellWidget(r, 0)
                cells.append("1" if (cw is not None and cw.isChecked()) else "0")
                continue
            it = first.item(r, c)
            cells.append(norm(it.text()) if it else "")
        return ";".join(cells)

    put("dlg.structure.rows", "\n".join(srow(r) for r in range(min(3, first.rowCount()))))
    sd.deleteLater()

    # ---- 说明与关于 ----
    hd = m.HelpDialog(win)
    put("dlg.help.title", hd.windowTitle())
    put("dlg.help.want", str(hd.property("e2mWant") or ""))
    put("env.dlg.help.size", "%dx%d" % (hd.width(), hd.height()))
    lst = hd.findChild(QListWidget)
    put("dlg.help.toc", "|".join(lst.item(i).text() for i in range(lst.count())))
    hd.deleteLater()
    ad = m.AboutDialog(win)
    put("dlg.about.title", ad.windowTitle())
    put("dlg.about.want", str(ad.property("e2mWant") or ""))
    put("env.dlg.about.size", "%dx%d" % (ad.width(), ad.height()))
    captions = []
    for lay in ad.findChildren(QGridLayout):
        for i in range(lay.rowCount()):
            it = lay.itemAtPosition(i, 0)
            wd = it.widget() if it is not None else None
            if wd is not None and norm(wd.text()):
                captions.append(norm(wd.text()))
    put("dlg.about.rows", "|".join(captions))
    ad.deleteLater()

    # ---- 进度框文案 ----
    pd = m.ProgressDialog(u"正在解析文件…", win)
    put("dlg.progress.title", pd.windowTitle())
    put("dlg.progress.want", str(pd.property("e2mWant") or ""))
    put("env.dlg.progress.size", "%dx%d" % (pd.width(), pd.height()))
    pd.set_progress(1, 3, u"读盘")
    put("dlg.progress.status", pd.lbl_step.text())
    put("dlg.progress.percent", pd.bar.text())
    put("dlg.progress.cancel", pd.btn_cancel.text())
    pd.deleteLater()

    # ---- 生成标准/简易 excel：产物本身也进对撞 ----
    out_root = os.path.join(TMP, "out")
    os.makedirs(out_root)
    real_pdir = m.program_dir
    m.program_dir = lambda: out_root
    win._gen_excel(u"标准", True)
    app.processEvents()
    put("export.std.log", win.log.toPlainText())
    put("export.std.popup.title", next((t for _k, t, _b in POPUPS if t == u"生成结果"), ""))
    put("export.std.popup.text", popup_body(
        next((b for _k, _t, b in POPUPS if _t == u"生成结果"), "")))
    std = os.path.join(out_root, [x for x in sorted(os.listdir(out_root))
                                  if x.startswith("meta_data_file_")][0])
    put("export.std.dir", os.path.basename(std))
    put("export.std.files", "|".join(sorted(os.listdir(std))))
    put("export.std.content", dump_xlsx(std))
    POPUPS[:] = []
    win._gen_excel(u"简易", False)
    app.processEvents()
    sim = os.path.join(out_root, [x for x in sorted(os.listdir(out_root))
                                  if x.startswith("simple_meta_data_file_")][0])
    put("export.simple.dir", os.path.basename(sim))
    put("export.simple.content", dump_xlsx(sim))
    m.program_dir = real_pdir

    # ---- SQL 页 ----
    sql_dir = os.path.join(TMP, "sql")
    os.makedirs(sql_dir)
    a = os.path.join(sql_dir, "a.sql")
    with open(a, "w", encoding="utf-8") as fh:
        fh.write("CREATE TABLE t_order (id INT);\n"
                 "INSERT INTO `t_order` VALUES (1);\n"
                 "CREATE TABLE plain (id INT);\n")
    b = os.path.join(sql_dir, "b.sql")
    with open(b, "w", encoding="utf-8") as fh:
        fh.write("CREATE TABLE dim_x (id INT);\n")
    sink = m.ProgressSink(on_log=lambda t: m.append_log(win.sql_log, t),
                          on_log_error=lambda t: m.append_log(win.sql_log, t, error=True))
    win._finish_sql_load(m.sql_load_task([a, b], [], sink))
    app.processEvents()
    dump_table("sql", win.sql_table)
    put("sql.log", win.sql_log.toPlainText())
    put("sql.add.popup", popup_text().replace(sql_dir, u"根目录"))
    win.ed_sql_prefix.setText("ods_")
    win.on_sql_apply_prefix()
    app.processEvents()
    dump_table("sql", win.sql_table)
    put("sql.prefixed.log", win.sql_log.toPlainText())
    put("sql.hint", win.lbl_sql_db.text())
    POPUPS.clear()
    win.ed_sql_prefix.setText("")
    win.on_sql_apply_prefix()
    put("guard.sql.empty_prefix", popup_text())

    # ---- SQL 页：保存副本、没选库就执行、整份清空 ----
    # 「清空待上传文件」放在保存/执行之后取证，否则后两步没有脚本可处理，
    # 比的就是空清单守卫而不是保存与执行
    POPUPS.clear()
    smark = win.sql_log.toPlainText()
    win.on_sql_save_modified()
    app.processEvents()
    put("sql.save.log", win.sql_log.toPlainText()[len(smark):].strip())
    put("sql.save.popup", popup_text().replace(sql_dir, u"根目录"))
    POPUPS.clear()
    win.on_sql_execute()
    put("sql.execute.nodb", popup_text())
    cmark = win.sql_log.toPlainText()
    win.on_sql_clear_files()
    put("sql.remove.log", win.sql_log.toPlainText()[len(cmark):].strip())
    gdir = os.path.join(TMP, "sqlgen")
    os.makedirs(gdir)
    probe = os.path.join(gdir, "probe.sql")
    with open(probe, "w", encoding="utf-8") as fh:
        fh.write(u"CREATE TABLE `t_gen` (\n"
                 u"  `id` bigint(20) unsigned NOT NULL AUTO_INCREMENT COMMENT '主键',\n"
                 u"  `st` enum('a','b') DEFAULT 'a',\n"
                 u"  PRIMARY KEY (`id`)\n"
                 u") ENGINE=InnoDB COMMENT='生成检查表';\n"
                 u"CREATE TABLE t_dup9 (x INT);\n"
                 u"CREATE TABLE t_dup9 (y INT);\n"
                 u"CREATE TABLE t_like9 LIKE t_dup9;\n")
    gmark = win.sql_log.toPlainText()
    win._finish_sql_load(m.sql_load_task([probe], [], sink))
    app.processEvents()
    std_root = os.path.join(gdir, "std")
    smp_root = os.path.join(gdir, "smp")
    os.makedirs(std_root)
    os.makedirs(smp_root)
    POPUPS[:] = []
    m.program_dir = lambda: std_root
    win.on_sql_gen_excel(u"标准", True)
    app.processEvents()
    m.program_dir = lambda: smp_root
    win.on_sql_gen_excel(u"简易", False)
    app.processEvents()
    m.program_dir = real_pdir
    gen_log = win.sql_log.toPlainText()[len(gmark):].strip()
    gen_log = gen_log.replace(gdir, u"根目录")
    put("sqlgen.log", re.sub(r"(simple_)?meta_data_file_\d{8}",
                             r"\1meta_data_file_日期", gen_log))
    # 生成结果弹窗：标题与正文两边都要同一句（正文里的目录同样归一化）
    put("sqlgen.popup", re.sub(r"(simple_)?meta_data_file_\d{8}",
                               r"\1meta_data_file_日期",
                               popup_text().replace(gdir, u"根目录")))
    std = os.path.join(std_root, [x for x in sorted(os.listdir(std_root))
                                  if x.startswith("sql_meta_data_file_")][0])
    smp = os.path.join(smp_root, [x for x in sorted(os.listdir(smp_root))
                                  if x.startswith("sql_simple_meta_data_file_")][0])
    put("sqlgen.std.content", dump_xlsx(std))
    put("sqlgen.simple.content", dump_xlsx(smp))
    # 用本条线自己的读取器把标准件读回来，渲染成一行给对撞器比（读不回就是缺陷）
    try:
        gen_file = os.path.join(std, [x for x in sorted(os.listdir(std))
                                      if "t_gen" in x][0])
        td = m.build_table_def(gen_file)
        line = [td["source"], td["table_name"], td.get("table_comment") or ""]
        for c in td["columns"]:
            line.append("%s:%s:%s" % (c["field"], c["type"],
                                      "pk" if c.get("primary_key") else "-"))
        put("sqlgen.roundtrip", "|".join(line))
    except Exception as e:
        put("sqlgen.roundtrip", "读不回 " + type(e).__name__)

    win.on_clear()
    app.processEvents()
    dump_table("file", win.file_table)
    for t in ("开始导入", "生成标准excel文件", "生成简易excel文件"):
        put("cleared.on.%s" % t, btn_on(win, t))

    # ---- 三页顶部的目标库信息条必须同源：塞一个库进去，三句话和三种颜色要一起变 ----
    # 各页各写一份刷新的话，迟早会在某个页签看到"上一个库"；host:port 只在有库时才进文案，
    # 所以这里必须真选一个库，不能拿空态糊过去。放在最后一步做，免得把前面那些
    # "没选库当场拦下"的探针带偏（它们断的就是空态那句话）。
    win.cb_db.addItem(u"alpha", u"alpha")
    app.processEvents()
    put("hint.db.linked", u"|".join([win.lbl_imp_db.text(), win.lbl_sql_db.text(),
                                     win.lbl_run_db.text()]))
    put("hint.db.tone", u"|".join(_db_tone(w) for w in
                                  (win.lbl_imp_db, win.lbl_sql_db, win.lbl_run_db)))

    out_dir = os.path.dirname(os.path.abspath(OUT))
    if not os.path.isdir(out_dir):
        os.makedirs(out_dir)
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(D, fh, ensure_ascii=False, indent=1, sort_keys=True)
    print("+ [OK] GUI-DUMP-PY | %d 个键 | %s" % (len(D), OUT))
    shutil.rmtree(TMP, ignore_errors=True)
    return 0


def dump_xlsx(folder):
    """目录里每个 xlsx 的 Sheet 名与逐格值，压成一段可比对的文本。"""
    from openpyxl import load_workbook

    out = []
    for name in sorted(os.listdir(folder)):
        wb = load_workbook(os.path.join(folder, name), read_only=True, data_only=True)
        out.append("### " + name)
        for ws in wb.worksheets:
            out.append("[%s]" % ws.title)
            for row in ws.iter_rows(values_only=True):
                out.append(";".join("" if v is None else str(v) for v in row))
        wb.close()
    return "\n".join(out)


if __name__ == "__main__":
    sys.exit(main())
