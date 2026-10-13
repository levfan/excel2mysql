# -*- coding: utf-8 -*-
"""core/meta_export.py 的回读测试：生成的标准件必须能被 build_table_def 原样读回来。

为什么盯"回读"而不是比对生成的单元格：这个功能的产物是给人继续导入用的，
只要 build_table_def 读回来的表定义和原来一致，命名、列顺序、类型拆分、
主键与可空的写法就全都对上了；反过来，只比对新写出来的格子文本，
读不读得回来是查不出来的——而读不回来就等于这个按钮白做。

跑法：.venv/Scripts/python.exe test_meta_export.py [夹具目录]
"""
from __future__ import print_function

import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from core.infer import build_table_def          # noqa: E402
from core.merge_check import MergedCellsError   # noqa: E402
from core import meta_export as mx              # noqa: E402

FIX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "java", "verify", "fixtures")


def expectation(cols):
    """普通推断出的表没有可空声明；导出件把它显式化，口径与建表一致。"""
    out = []
    for c in cols:
        pk = bool(c.get("primary_key"))
        out.append((c["field"], c["type"], c["comment"], pk, not pk))
    return out


def read_back(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


class _Sink(object):
    """只记账不渲染的回报口：写够几个就 pretend 用户按了停止。"""

    def __init__(self, stop_after):
        self.stop_after = stop_after
        self.done = 0
        self.cancelled = False
        self.total = -1
        self.lines = []

    def set_total(self, n):
        self.total = n

    def log(self, text):
        self.lines.append(text)

    def log_error(self, text):
        self.lines.append(text)

    def step(self, name):
        self.done += 1
        if self.done >= self.stop_after:
            self.cancelled = True


def cancel_stops_the_batch():
    """按了停止之后剩下的文件不能再写：收尾那句"已取消，剩余文件未生成"得是真话。"""
    out = []
    work = tempfile.mkdtemp(prefix="e2m-meta-cancel-")
    try:
        tables = []
        for name in sorted(os.listdir(FIX)):
            if not name.lower().endswith((".xlsx", ".xls")):
                continue
            try:
                tables.append(build_table_def(os.path.join(FIX, name)))
            except MergedCellsError:
                continue
            if len(tables) >= 3:
                break
        if len(tables) < 2:
            return ["用来验取消的夹具不足 2 个"]
        sink = _Sink(1)
        _target, written = mx.export_tables(tables, work, True, sink)
        if len(written) != 1:
            out.append("取消后还在写文件：写了 %d 个，应为 1 个" % len(written))
        if sink.total != len(tables):
            out.append("总项数没在开跑前算准：%s / %d" % (sink.total, len(tables)))
        if not any(line.startswith("+ [OK] 导出 | ") for line in sink.lines):
            out.append("写完一个文件没有留下成功行")
    finally:
        shutil.rmtree(work, ignore_errors=True)
    return out


def main():
    files = sorted(f for f in os.listdir(FIX)
                   if f.lower().endswith((".xlsx", ".xls")))
    problems = []
    checked = 0
    skipped = 0
    work = tempfile.mkdtemp(prefix="e2m-meta-export-")
    try:
        for name in files:
            src = os.path.join(FIX, name)
            try:
                table = build_table_def(src)
            except MergedCellsError:
                skipped += 1     # 含合并单元的夹具本来就该被拒，不是导出的问题
                continue
            checked += 1

            # 标准件：元数据 + 数据，且能被读回来
            std_dir = os.path.join(work, "std", name)
            os.makedirs(std_dir)
            _target, written = mx.export_tables([table], std_dir, True)
            if len(written) != 1:
                problems.append("%s 标准件没写成文件" % name)
                continue
            std = written[0]
            expect_name = mx.file_name(table)
            if os.path.basename(std) != expect_name:
                problems.append("命名不对：%s，应为 %s" % (os.path.basename(std), expect_name))
            sheets = read_back(std)
            if sheets != [mx.META_SHEET, mx.DATA_SHEET]:
                problems.append("标准件 Sheet 应为 元数据+数据，实际 %s" % sheets)
            back = build_table_def(std)
            if back["source"] != "structured":
                problems.append("%s 回读没走结构定义那条路（source=%s）"
                                % (name, back["source"]))
            for key in ("table_name", "table_comment", "primary_key"):
                if table.get(key) != back.get(key):
                    problems.append("%s 回读后 %s 变了：%r -> %r"
                                    % (name, key, table.get(key), back.get(key)))
            got = [(c["field"], c["type"], c["comment"], bool(c.get("primary_key")),
                    bool(c.get("nullable"))) for c in back["columns"]]
            want = expectation(table["columns"])
            if got != want:
                first = next((i for i, (a, b) in enumerate(zip(got, want)) if a != b), None)
                problems.append("%s 第 %s 列回读不一致：%r -> %r"
                                % (name, first, want[first] if first is not None else None,
                                   got[first] if first is not None else None))

            # 简易件：只有数据一个 Sheet
            sim_dir = os.path.join(work, "sim", name)
            os.makedirs(sim_dir)
            _t2, w2 = mx.export_tables([table], sim_dir, False)
            if len(w2) != 1:
                problems.append("%s 简易件没写成文件" % name)
                continue
            sim_sheets = read_back(w2[0])
            if sim_sheets != [mx.DATA_SHEET]:
                problems.append("%s 简易件应只有数据 Sheet，实际 %s" % (name, sim_sheets))
            sim_head = []
            import openpyxl
            wb = openpyxl.load_workbook(w2[0], read_only=True)
            for row in wb[mx.DATA_SHEET].iter_rows(max_row=1, values_only=True):
                sim_head = list(row)
            wb.close()
            if sim_head != [c["field"] for c in table["columns"]]:
                problems.append("%s 简易件数据表头不是字段名：%s" % (name, sim_head))

            # 行数要对得上：数据一格不少
            import pandas as pd
            src_rows = len(table["df"].index) if table.get("df") is not None else 0
            back_rows = len(back["df"].index) if back.get("df") is not None else 0
            if src_rows != back_rows:
                problems.append("%s 回读后行数变了：%d -> %d" % (name, src_rows, back_rows))
    finally:
        shutil.rmtree(work, ignore_errors=True)

    problems += cancel_stops_the_batch()

    print("- [INFO] META-EXPORT | 回读比对 %d 个夹具（跳过含合并单元 %d 个）"
          % (checked, skipped))
    for p in problems:
        print("x [ERROR] META-EXPORT | " + p)
    if problems:
        print("x [ERROR] META-EXPORT | 导出与回读不一致")
        return 1
    print("+ [OK] META-EXPORT | 标准件可原样读回，简易件只剩数据 Sheet，行数与列定义一致")
    return 0


if __name__ == "__main__":
    sys.exit(main())
